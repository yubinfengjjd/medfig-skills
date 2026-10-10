"""medfig-outline scripts: outline_check (numbers / references / wording) and outline_build (docx + md + folders).
Synthetic mini-project only."""
import json
import sys
from pathlib import Path

import pytest
import yaml

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import outline_build  # noqa: E402
import outline_check  # noqa: E402


@pytest.fixture
def proj(tmp_path):
    """Mini project: 2 figures with source.json + pdf/png, 1 table, captions, scripts."""
    p = tmp_path / "proj"
    (p / "out/figures/main").mkdir(parents=True)
    (p / "out/tables").mkdir(parents=True)
    (p / "docs/captions").mkdir(parents=True)
    (p / "figures/main").mkdir(parents=True)
    (p / "tables").mkdir()
    for name, vals in {"fig2": {"auc": {"all": 0.98312}, "n": 13859, "panels": [{"id": "a"}, {"id": "b"}]},
                       "fig3": {"recall": {"Edema": 0.9273}, "panels": [{"id": "a"}]}}.items():
        src = {"figure": name, "inputs": [{"path": "data:mainline/tables/x.csv", "sha256": "0" * 64, "rows": 3}],
               "transforms": [], "values": vals}
        (p / f"out/figures/main/{name}.source.json").write_text(json.dumps(src), encoding="utf-8")
        for ext in ("pdf", "png", "svg"):
            (p / f"out/figures/main/{name}.{ext}").write_bytes(b"x")
        (p / f"figures/main/{name}.py").write_text("# script", encoding="utf-8")
    (p / "out/tables/T1.csv").write_text("cohort,unit,bacc_mean\nInternal test,scan,0.9623\nSite B,scan,0.929\n",
                                         encoding="utf-8")
    (p / "out/tables/T1.md").write_text("**T1.** x\n", encoding="utf-8")
    (p / "tables/build_tables.py").write_text("# tables", encoding="utf-8")
    (p / "src").mkdir()
    for name in ("splits.py", "eval.py"):
        (p / "src" / name).write_text(f"# {name}", encoding="utf-8")
    (p / "docs/captions/fig2.md").write_text("## Fig2\n\n**图 2** x\n", encoding="utf-8")
    (p / "docs/captions_zh.md").write_text("## Fig3\n\n**图 3** y\n", encoding="utf-8")
    return p


def _outline(**over):
    o = {
        "project": "demo", "figures_root": "out/figures", "tables_root": "out/tables",
        "results_fewer_sections_reason": "synthetic two-figure project",
        "captions": ["docs/captions_zh.md", "docs/captions"], "main_figures": ["fig2", "fig3"],
        "methods": [{"id": "2.1", "title": "Cohorts and evaluation", "short": "Cohorts",
                     "focus": [{"topic": "研究设计", "points": ["多中心回顾性研究；三分类（内部代号 E2）", "说明外部队列不参与训练"]}],
                     "items": [{"ref": "T1", "what": "各队列规模与角色"}, {"ref": "Fig1", "what": "研究设计与数据流"}],
                     "code": [{"path": "src/splits.py", "role": "计算", "guess": True}]}],
        "design_assets": ["T1", "Fig1"], "pending": ["Fig1"],
        "results": [
            {"id": "3.1", "title": "Cohorts separate in the backbone representation", "subtitle": "队列在骨干表示中分开",
             "short": "Domain_shift",
             "paragraphs": [{"label": "第一段", "claim": "骨干层面队列可区分",
                             "points": ["域可辨识性 AUC 约 0.983，n = 13,859", "说明 PCA 只在内部拟合"], "cites": ["Fig2a-b", "T1"]}],
             "items": [{"ref": "Fig2", "what": "外部队列在骨干表示中各自成簇"}], "boundaries": ["不写设备因果"],
             "code": [{"path": "src/eval.py", "role": "计算", "guess": True}]},
            {"id": "3.2", "title": "Edema recall is high internally", "subtitle": "内部 Edema 召回率高",
             "short": "Internal", "paragraphs": [{"label": "第一段", "claim": "召回高",
                                                   "points": ["Edema 召回率 92.7%；内部 BACC 约 0.96", "说明主要错误方向"],
                                                   "cites": ["Fig3a", "T1"]}],
             "items": [{"ref": "Fig3a", "what": "三类召回都高"}], "code": []}],
        "numbers": [{"text": "0.983", "source": "fig2.source.json:values.auc.all"},
                    {"text": "13,859", "source": "fig2.source.json:values.n"},
                    {"text": "92.7%", "source": "fig3.source.json:values.recall.Edema"},
                    {"text": "0.96", "source": "table:T1.csv:bacc_mean@cohort=Internal test&unit=scan"}],
    }
    o.update(over)
    return o


def _write(tmp_path, o):
    f = tmp_path / "outline.yaml"
    f.write_text(yaml.safe_dump(o, allow_unicode=True, sort_keys=False), encoding="utf-8")
    return f


def test_clean_outline_passes(proj, tmp_path):
    assert outline_check.check(_write(tmp_path, _outline()), proj) == []


def test_wrong_number_and_unregistered_number(proj, tmp_path):
    o = _outline()
    o["numbers"][0]["text"] = "0.984"
    o["results"][0]["paragraphs"][0]["points"].append("另一个数 0.55")
    out = outline_check.check(_write(tmp_path, o), proj)
    assert any("0.984" in i and "source value 0.98312" in i for i in out)
    assert any("0.55" in i and "not registered" in i for i in out)


def test_missing_refs_and_uncovered_main_figure(proj, tmp_path):
    o = _outline()
    o["results"][0]["paragraphs"][0]["cites"] = ["Fig2c", "ST99"]
    o["results"][1]["paragraphs"][0]["cites"] = ["T1"]
    o["results"][1]["items"] = []
    out = outline_check.check(_write(tmp_path, o), proj)
    assert any("Fig2c" in i for i in out) and any("ST99" in i for i in out)
    assert any("fig3" in i and "not covered" in i for i in out)


def test_methods_focus_has_no_result_numbers(proj, tmp_path):
    o = _outline()
    o["methods"][0]["focus"][0]["points"].append("AUC 达到 0.98")
    out = outline_check.check(_write(tmp_path, o), proj)
    assert any("Methods 2.1" in i and "0.98" in i for i in out)
    o["methods"][0]["allow_numbers"] = ["0.98"]
    assert not any("Methods 2.1" in i for i in outline_check.check(_write(tmp_path, o), proj))


@pytest.mark.parametrize("text", ["与早期版本相比", "pre-registered hypotheses", "Wave-1 model", "closure analysis",
                                  "预注册假设", "修复后的模型", "validated performance"])
def test_wording_gates(proj, tmp_path, text):
    o = _outline()
    o["results"][0]["paragraphs"][0]["points"].append(text)
    out = outline_check.check(_write(tmp_path, o), proj)
    assert out, text


def test_internal_code_only_inside_code_note(proj, tmp_path):
    toml = tmp_path / "figkit.toml"
    toml.write_text("data_root = \"d\"\n[qa]\nforbidden_patterns = ['\\bE[12]\\b']\n", encoding="utf-8")
    assert outline_check.check(_write(tmp_path, _outline()), proj, figkit_toml=toml) == []  # "（内部代号 E2）" ok
    o = _outline()
    o["results"][0]["paragraphs"][0]["points"].append("E2 头的召回")
    out = outline_check.check(_write(tmp_path, o), proj, figkit_toml=toml)
    assert any("E[12]" in i for i in out)


def test_results_title_must_be_a_finding(proj, tmp_path):
    o = _outline()
    o["results"][0]["title"] = "Results of domain shift"
    o["results"][1]["title"] = "AUC was 0.98 internally"
    out = outline_check.check(_write(tmp_path, o), proj)
    assert any("3.1" in i and "finding" in i for i in out)
    assert any("3.2" in i and "number" in i for i in out)


def test_build_docx_md_and_folders(proj, tmp_path):
    out = tmp_path / "outline_out"
    res = outline_build.build(_write(tmp_path, _outline()), proj, out)
    import docx
    d = docx.Document(res["docx"])
    text = "\n".join(p.text for p in d.paragraphs)
    for s in ("2.1 Cohorts and evaluation", "3.1 Cohorts separate in the backbone representation",
              "（队列在骨干表示中分开）", "写作重点", "写作要点", "第一段", "图 / 表", "相关代码", "措辞边界",
              "src/eval.py（推测）"):
        assert s in text, s
    md = Path(res["md"]).read_text(encoding="utf-8")
    assert "## 3.1 Cohorts separate" in md
    r31 = out / "3_Results/3.1_Domain_shift"
    assert (r31 / "figures/fig2.pdf").is_file() and (r31 / "figures/fig2.source.json").is_file()
    assert (r31 / "tables/T1.csv").is_file() and (r31 / "tables/T1.md").is_file()
    assert (r31 / "captions/fig2.md").is_file()
    assert not (r31 / "code/fig2.py").exists()  # plotting scripts are not part of the method
    assert (r31 / "code/eval.py").read_text(encoding="utf-8") == "# eval.py"  # source file, guess or not
    assert not (r31 / "code/code.md").exists()  # the path list lives in the outline, not in the folder
    assert (out / "2_Methods/2.1_Cohorts/code/splits.py").is_file()
    assert (out / "2_Methods/2.1_Cohorts/tables/T1.csv").is_file()
    cap3 = (out / "3_Results/3.2_Internal/captions/fig3.md").read_text(encoding="utf-8")
    assert "图 3" in cap3  # caption section extracted from the combined captions file
    assert (proj / "out/figures/main/fig2.pdf").is_file()  # copies, not moves


def test_code_paths_must_exist_and_code_roots_resolve(proj, tmp_path):
    o = _outline()
    o["results"][0]["code"].append({"path": "repo/tools/run.py", "role": "计算", "guess": True})
    o["results"][1]["code"].append({"path": "TODO", "role": "计算"})
    out = outline_check.check(_write(tmp_path, o), proj)
    assert out == [i for i in out if "repo/tools/run.py" in i and "code_roots" in i] and len(out) == 1
    upstream = tmp_path / "upstream"
    (upstream / "repo/tools").mkdir(parents=True)
    (upstream / "repo/tools/run.py").write_text("# run", encoding="utf-8")
    o["code_roots"] = [str(upstream)]
    assert outline_check.check(_write(tmp_path, o), proj) == []
    out_dir = tmp_path / "outline_out"
    outline_build.build(_write(tmp_path, o), proj, out_dir)
    assert (out_dir / "3_Results/3.1_Domain_shift/code/run.py").read_text(encoding="utf-8") == "# run"


def test_same_name_code_files_keep_their_paths(proj, tmp_path):
    (proj / "src/a").mkdir()
    (proj / "src/a/eval.py").write_text("# a/eval.py", encoding="utf-8")
    o = _outline()
    o["results"][0]["code"].append({"path": "src/a/eval.py", "role": "计算"})
    out_dir = tmp_path / "outline_out"
    outline_build.build(_write(tmp_path, o), proj, out_dir)
    code = out_dir / "3_Results/3.1_Domain_shift/code"
    assert (code / "src/eval.py").read_text(encoding="utf-8") == "# eval.py"
    assert (code / "src/a/eval.py").read_text(encoding="utf-8") == "# a/eval.py"
    assert not (code / "eval.py").exists()


def test_build_refuses_output_inside_project(proj, tmp_path):
    with pytest.raises(ValueError, match="outside the project"):
        outline_build.build(_write(tmp_path, _outline()), proj, proj / "outline")


def test_build_refuses_failing_outline(proj, tmp_path):
    o = _outline()
    o["numbers"][0]["text"] = "0.984"
    with pytest.raises(RuntimeError, match="outline_check"):
        outline_build.build(_write(tmp_path, o), proj, tmp_path / "o2")


def test_number_forms():
    m = outline_check._matches
    assert m("1/3", 0.3333333333) and not m("1/3", 0.34)
    assert m("+0.017", 0.017457) and not m("+0.017", -0.017457) and m("−0.019", -0.019026)
    assert m("92.7%", 0.927318) and m("13,859", 13859)
    assert m("0.047", -0.0472, magnitude=True) and not m("0.047", -0.0472)  # "低约 0.047" = magnitude
    assert outline_check.NUM.findall("P = 1/3，差值 +0.00127") == ["1/3", "+0.00127"]


def test_style_limits(proj, tmp_path):
    o = _outline()
    p = o["results"][0]["paragraphs"][0]
    p["points"] = ["点出 AUC 约 0.983", "交代 n = 13,859", "说明另一点", "第四条要点"]
    out = outline_check.check(_write(tmp_path, o), proj)
    assert any("4 points (max 3)" in i for i in out)
    p["points"] = ["引用 Fig 2a：说明骨干层可分"]
    out = outline_check.check(_write(tmp_path, o), proj)
    assert any("opens with a citation" in i for i in out)
    p["points"] = ["说明" + "很长的要点" * 20]
    out = outline_check.check(_write(tmp_path, o), proj)
    assert any("chars (max 70" in i for i in out)
    o["style"] = {"point_max_chars": 200}
    assert not any("chars (max" in i for i in outline_check.check(_write(tmp_path, o), proj))


def test_style_numbers_per_point_and_paragraph(proj, tmp_path):
    o = _outline()
    o["numbers"] += [{"text": "0.98", "source": "fig2.source.json:values.auc.all"}]
    o["results"][0]["paragraphs"][0]["points"] = ["对比 0.983、0.98 与 13,859"]
    out = outline_check.check(_write(tmp_path, o), proj)
    assert any("3 numbers in one point" in i for i in out)


def test_build_puts_citation_on_the_paragraph_line(proj, tmp_path):
    o = _outline()
    res = outline_build.build(_write(tmp_path, o), proj, tmp_path / "o3")
    import docx
    text = [p.text for p in docx.Document(res["docx"]).paragraphs]
    assert "第一段：骨干层面队列可区分（Fig 2a-b、Table 1）" in text
    assert not any(t.startswith("引用：") for t in text)
    assert "计算：src/eval.py（推测）" in text and not any("绘图" in t for t in text)
    assert "Fig 1（外部绘制，待交付）：研究设计与数据流" in text


def test_style_range_counts_as_one_number(proj, tmp_path):
    o = _outline()
    o["numbers"] += [{"text": "0.98", "source": "fig2.source.json:values.auc.all"}]
    o["results"][0]["paragraphs"][0]["points"] = ["点出范围 0.98–0.983，n = 13,859"]
    assert not any("numbers in one point" in i for i in outline_check.check(_write(tmp_path, o), proj))



def _nested():
    o = _outline()
    m = o["methods"][0]
    m["subsections"] = [
        {"id": "2.1.1", "title": "Stage 1", "short": "Stage1", "focus": [{"topic": "前端", "points": ["说明边界", "说明软带"]}]},
        {"id": "2.1.2", "title": "Stage 2", "short": "Stage2", "focus": [{"topic": "联合训练", "points": ["说明损失", "说明冻结"]}]}]
    r = o["results"][0]
    r["subsections"] = [
        {"id": "3.1.1", "title": "Cohort composition", "short": "Cohorts", "subtitle": "队列构成",
         "paragraphs": [{"label": "第一段", "claim": "构成不同", "points": ["点出构成差异", "说明两类队列"], "cites": ["T1"]}]},
        {"id": "3.1.2", "title": "Representation", "short": "Repr", "subtitle": "表示",
         "paragraphs": [{"label": "第一段", "claim": "可分", "points": ["说明 PCA", "对比空间"], "cites": ["Fig2a"]}]}]
    return o


def test_subsections_are_checked_and_rendered(proj, tmp_path):
    o = _nested()
    assert outline_check.check(_write(tmp_path, o), proj) == []
    o["results"][0]["subsections"][1]["paragraphs"][0]["cites"] = ["Fig2z"]
    assert any("3.1.2" in i and "Fig2z" in i for i in outline_check.check(_write(tmp_path, o), proj))
    o = _nested()
    res = outline_build.build(_write(tmp_path, o), proj, tmp_path / "o4")
    import docx
    paras = docx.Document(res["docx"]).paragraphs
    heads = {p.text: p.runs[0].font.size.pt for p in paras if p.runs and p.runs[0].bold and p.runs[0].font.size}
    assert heads["2.1 Cohorts and evaluation"] == 14 and heads["2.1.1 Stage 1"] == 12
    assert heads["3.1.2 Representation"] == 12
    out = tmp_path / "o4"
    assert (out / "3_Results/3.1_Domain_shift/3.1.1_Cohorts/tables/T1.csv").is_file()
    assert (out / "3_Results/3.1_Domain_shift/3.1.2_Repr/figures/fig2.pdf").is_file()
    assert "### 2.1.1 Stage 1" in Path(res["md"]).read_text(encoding="utf-8")


def test_structure_rules(proj, tmp_path):
    o = _nested()
    o["methods"][0]["subsections"] = o["methods"][0]["subsections"][:1]
    assert any("single subsection" in i for i in outline_check.check(_write(tmp_path, o), proj))
    o = _nested()
    o["results"][0]["subsections"][0]["id"] = "3.2.1"
    assert any("not numbered under 3.1" in i for i in outline_check.check(_write(tmp_path, o), proj))
    o = _nested()
    o["results"][0]["subsections"][0]["paragraphs"][0]["points"] = ["只有一条"]
    assert any("3.1.1 has 1 point" in i for i in outline_check.check(_write(tmp_path, o), proj))


def test_mechanical_one_figure_per_section_is_flagged(proj, tmp_path):
    o = _outline()
    sec = o["results"][1]
    o["results"] = [dict(sec, id=f"3.{k}", short=f"S{k}") for k in range(1, 5)]
    o["methods"] = [dict(o["methods"][0], id=f"2.{k}", short=f"M{k}") for k in range(1, 5)]
    o["main_figures"] = ["fig3"]
    out = outline_check.check(_write(tmp_path, o), proj)
    assert any("same number of sections" in i for i in out)


def test_results_section_count_range(proj, tmp_path):
    o = _outline()
    del o["results_fewer_sections_reason"]
    out = outline_check.check(_write(tmp_path, o), proj)
    assert any("only 2 top-level Results sections (aim for 5-8)" in i for i in out)
    o["results_fewer_sections_reason"] = "两步论证：先性能后机制"
    assert not any("top-level Results sections" in i for i in outline_check.check(_write(tmp_path, o), proj))
    sec = o["results"][1]
    o["results"] = [dict(sec, id=f"3.{k}", short=f"S{k}") for k in range(1, 10)]
    assert any("9 top-level Results sections (max 8)" in i for i in outline_check.check(_write(tmp_path, o), proj))


# ---- citation order (outline_format.md section 7)
def _issues(proj, tmp_path, o):
    return outline_check.check(_write(tmp_path, o), proj)


def test_panels_first_cited_in_letter_order(proj, tmp_path):
    o = _outline()
    o["results"][0]["paragraphs"][0]["cites"] = ["Fig2b", "T1"]
    o["results"][0]["paragraphs"][0]["points"][1] = "说明 PCA 只在内部拟合（Fig 2a）"
    out = _issues(proj, tmp_path, o)
    assert any("first cites Fig 2b before Fig 2a" in i for i in out)
    o["results"][0]["paragraphs"][0]["cites"] = ["Fig2a", "T1"]
    o["results"][0]["paragraphs"][0]["points"][1] = "说明 PCA 只在内部拟合（Fig 2b）"
    assert not any("order:" in i for i in _issues(proj, tmp_path, o))


def test_figures_first_cited_in_number_order_with_renumber_plan(proj, tmp_path):
    o = _outline()
    o["results"][0]["paragraphs"][0]["cites"] = ["Fig3a", "Fig2a-b", "T1"]
    out = _issues(proj, tmp_path, o)
    hit = [i for i in out if i.startswith("order: main figs")]
    assert hit and "Fig 1, Fig 3, Fig 2" in hit[0] and "Fig 3 → Fig 2" in hit[0] and "Fig 2 → Fig 3" in hit[0]


def test_text_mentions_count_as_citations(proj, tmp_path):
    o = _outline()
    o["results"][0]["paragraphs"][0]["cites"] = ["T1"]
    o["results"][0]["paragraphs"][0]["points"][1] = "对比 Fig 3a 的召回与 Fig 2a–b 的分布"
    assert any("main figs are first cited as Fig 1, Fig 3, Fig 2" in i for i in _issues(proj, tmp_path, o))


def test_supp_series_uncited_assets_and_excluded(proj, tmp_path):
    (proj / "out/figures/supp").mkdir()
    for name in ("s01", "s02"):
        (proj / f"out/figures/supp/{name}.source.json").write_text(
            json.dumps({"values": {"panels": [{"id": "a"}, {"id": "b"}]}}), encoding="utf-8")
    (proj / "out/tables/ST01.csv").write_text("x\n1\n", encoding="utf-8")
    o = _outline()
    out = _issues(proj, tmp_path, o)
    assert any("Fig S1 is exported but never cited" in i for i in out)
    assert any("Table S1 is exported but never cited" in i for i in out)
    p = o["results"][1]["paragraphs"][0]
    p["cites"] = ["Fig3a", "S2", "S1a", "T1"]
    o["results"][1]["items"] += [{"ref": "S1", "what": "重复处理规则"}, {"ref": "S2", "what": "逐 seed 稳定"}]
    o["excluded"] = [{"ref": "ST01", "reason": "溯源索引，不进稿件"}]
    out = _issues(proj, tmp_path, o)
    assert any("supp figs are first cited as Fig S2, Fig S1" in i for i in out)
    assert any("Fig S1 panel(s) b never cited" in i for i in out)
    assert not any("Table S1" in i for i in out)
    o["excluded"] = ["ST01"]
    assert any("excluded 'ST01': give a reason" in i for i in _issues(proj, tmp_path, o))


def test_methods_cite_only_design_assets(proj, tmp_path):
    o = _outline()
    o["methods"][0]["focus"][0]["points"][1] = "说明外部队列不参与训练，召回见 Fig 3a"
    out = _issues(proj, tmp_path, o)
    assert any("Methods 2.1 cites Fig 3; Methods may cite only design_assets" in i for i in out)
    o = _outline(design_assets=["T1"])
    assert any("Methods 2.1 cites Fig 1" in i for i in _issues(proj, tmp_path, o))


def test_items_say_what_and_code_has_no_plot_scripts(proj, tmp_path):
    o = _outline()
    o["results"][1]["items"] = []
    o["results"][0]["items"] = [{"ref": "Fig2", "what": "由 fig2.py 绘制"}, {"ref": "Fig3", "what": "x"}]
    o["results"][0]["code"].append({"path": "figures/main/fig2.py", "role": "绘图"})
    o["methods"][0]["code"].append({"path": "tables/build_tables.py", "role": "表格"})
    out = _issues(proj, tmp_path, o)
    assert any("Results 3.1 Fig2: say what it shows scientifically" in i for i in out)
    assert any("Results 3.1 lists Fig 3 but its paragraphs do not cite it" in i for i in out)
    assert any("code: Results 3.1 lists figures/main/fig2.py" in i for i in out)
    assert any("code: Methods 2.1 lists tables/build_tables.py" in i for i in out)


def test_coverage_table_lists_every_asset(proj, tmp_path):
    (proj / "out/figures/supp").mkdir()
    (proj / "out/figures/supp/s01.source.json").write_text(json.dumps({"values": {"panels": [{"id": "a"}]}}),
                                                           encoding="utf-8")
    (proj / "out/tables/ST01.csv").write_text("x\n1\n", encoding="utf-8")
    o = _outline(excluded=[{"ref": "ST01", "reason": "溯源索引"}])
    rows = outline_check.coverage(o, proj)
    st = {r[0]: r[4] for r in rows}
    assert st == {"Fig 1": "pending", "Fig 2": "cited", "Fig 3": "cited", "Table 1": "cited",
                  "Fig S1": "NOT CITED", "Table S1": "excluded: 溯源索引"}
    assert [r[0] for r in rows] == ["Fig 1", "Fig 2", "Fig 3", "Table 1", "Fig S1", "Table S1"]
    assert any("Fig S1 is exported but never cited" in i for i in outline_check.check(_write(tmp_path, o), proj))


def test_build_writes_coverage_table(proj, tmp_path):
    res = outline_build.build(_write(tmp_path, _outline()), proj, tmp_path / "o5")
    md = Path(res["md"]).read_text(encoding="utf-8")
    assert "## 图表引用清单" in md and "| Fig 2 | a,b | 3.1 | 3.1 | 已引用 |" in md
    assert "| Fig 1 | — | 2.1 | 2.1 | 已引用（外部绘制，待交付） |" in md
    import docx
    t = docx.Document(res["docx"]).tables[-1]
    assert t.rows[0].cells[4].text == "状态" and len(t.rows) == 5


def test_figure_without_panel_list_is_flagged(proj, tmp_path):
    src = proj / "out/figures/main/fig3.source.json"
    d = json.loads(src.read_text(encoding="utf-8")); d["values"].pop("panels")
    src.write_text(json.dumps(d), encoding="utf-8")
    o = _outline()
    o["results"][1]["paragraphs"][0]["cites"] = ["Fig3", "T1"]
    o["results"][1]["items"] = [{"ref": "Fig3", "what": "三类召回都高"}]
    assert any("Fig 3 has no panel list" in i for i in outline_check.check(_write(tmp_path, o), proj))


def _with_case_selection(proj, seed=20261008):
    src = proj / "out/figures/main/fig3.source.json"
    d = json.loads(src.read_text(encoding="utf-8"))
    d["values"]["case_selection"] = {"seed": seed, "rule": "stratified random", "strata": ["truth", "correct"],
                                     "per_group": {"Normal": {"correct": 9, "misclassified": 3}},
                                     "pool": "available subset", "pool_counts": {"Normal|correct": 31}}
    src.write_text(json.dumps(d), encoding="utf-8")


def test_case_selection_must_be_stated_in_methods(proj, tmp_path):
    _with_case_selection(proj)
    out = outline_check.check(_write(tmp_path, _outline()), proj)
    assert any("case selection rule for Fig 3 not stated in Methods (seed 20261008)" in i for i in out)
    o = _outline()
    o["methods"][0]["focus"].append({"topic": "病例展示的选例", "points": [
        "按真实类别分层、固定 seed 20261008 随机抽取，每组 12 例（误判 3 例、正确 9 例），不按置信度挑选"]})
    assert outline_check.check(_write(tmp_path, o), proj) == []  # seed and per-group counts are design constants
    o["methods"][0]["focus"][-1]["points"] = ["按类别分层抽取病例，seed 20261008"]  # seed but no "random"
    assert any("case selection rule for Fig 3" in i for i in outline_check.check(_write(tmp_path, o), proj))


def test_case_selection_only_for_cited_figures(proj, tmp_path):
    _with_case_selection(proj)
    o = _outline(main_figures=["fig2"])
    o["results"] = o["results"][:1]
    out = outline_check.check(_write(tmp_path, o), proj)
    assert not any("case selection" in i for i in out)
