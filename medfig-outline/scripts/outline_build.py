"""Render a checked ``outline.yaml`` into 写作大纲.docx + 写作大纲.md and copy each subsection's files into its own
folder (figures / tables / captions / code), outside the project. Refuses to build when outline_check reports
issues. Only copies; never modifies the project.

Usage: python outline_build.py outline.yaml --project <root> --out <dir outside root> [--figkit-toml <toml>]
"""
import argparse
import re
import shutil
import sys
from pathlib import Path

import yaml

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import outline_check  # noqa: E402

FIG_EXT = ("pdf", "png", "svg")
TAB_EXT = ("csv", "md", "tex")
ROLE_ORDER = ["计算", "图内计算", "配置", "回执"]


# ------------------------------------------------------------------------------------------------ model
def _refs(sec, results):
    refs = [i["ref"] if isinstance(i, dict) else i for i in sec.get("items", [])] + list(sec.get("cites", []))
    for p in sec.get("paragraphs", []):
        refs += p.get("cites", [])
    figs, tabs = [], []
    for r in refs:
        for name, letter in outline_check._expand(r) or []:
            (tabs if name == "table" else figs).append(letter if name == "table" else name)
    return list(dict.fromkeys(figs)), list(dict.fromkeys(tabs))


def _items(sec):
    return [i for i in sec.get("items", []) if isinstance(i, dict)]


def _item_line(i, pending):
    tag = "（外部绘制，待交付）" if outline_check._asset(*(outline_check._expand(i["ref"]) or [("", "")])[0]) in pending else ""
    return f"{_pretty_ref(i['ref'])}{tag}：{i.get('what', '')}"


def _label(p):
    return p["label"] + ("（可选）" if p.get("optional") else "")


def _pretty_ref(r):
    """'Fig4a' -> 'Fig 4a', 'S5a' -> 'Fig S5a', 'ST05' -> 'Table S5', 'T2' -> 'Table 2'."""
    m = re.match(r"^(Fig|S)(\d+)(.*)$", r)
    if m:
        return f"Fig {'S' if m.group(1) == 'S' else ''}{int(m.group(2))}{m.group(3)}"
    m = re.match(r"^(ST|T)(\d+)$", r)
    if m:
        return f"Table {'S' if m.group(1) == 'ST' else ''}{int(m.group(2))}"
    return r


def _cite_tail(p):
    return f"（{'、'.join(_pretty_ref(c) for c in p['cites'])}）" if p.get("cites") else ""


def _code_roles(code):
    """[(role, line)] in 计算 / 图内计算 / 配置 / 回执 order. An entry with ``symbols`` / ``note`` gets its own line
    ('`decision.py` · `DecisionLayer.forward` —— 线性决策头'), so each point maps to the function that implements
    it; plain entries share one line per role. 配置 / 回执 lines say how the file is archived (摘录 / 只列路径)."""
    roles = ROLE_ORDER + sorted({c.get("role", "") for c in code} - set(ROLE_ORDER))
    out = []
    for r in roles:
        mine = [c for c in code if c.get("role", "") == r]
        plain = [c for c in mine if not (c.get("symbols") or c.get("note"))]
        for c in mine:
            if c in plain:
                continue
            syms = "、".join(f"`{s}`" for s in c.get("symbols") or [])
            out.append((r, f"`{_code_line(c)}`" + (f" · {syms}" if syms else "")
                        + (f" —— {c['note']}" if c.get("note") else "") + COPY_TAG.get(c.get("copy"), "")))
        if plain:
            out.append((r, "、".join(f"`{_code_line(c)}`{COPY_TAG.get(c.get('copy'), '')}" for c in plain)))
    return out


COPY_TAG = {"excerpt": "（摘录）", "path": "（只列路径，未拷贝）"}


def _methods_line(sec):
    return "方法：见 " + "、".join(str(m) for m in sec.get("methods") or [])


def _code_line(c):
    path = c.get("path") or "TODO"
    return f"{path}（推测）" if c.get("guess") else path


# ------------------------------------------------------------------------------------------------ markdown
HEAD_PT = {0: 14, 1: 12, 2: 11}   # docx heading size by depth (section / subsection / sub-subsection)


def _walk(secs):
    yield from outline_check._tree(secs)


def _md_body(sec, results, pending=frozenset()):
    L = []
    if not results:
        if sec.get("focus"):
            L += ["**写作重点：**", ""]
            for blk in sec["focus"]:
                L.append(f"- {blk['topic']}：")
                L += [f"  - {t}" for t in blk.get("points", [])]
        if sec.get("foreshadow"):
            L += ["", "**为后文铺垫：**", ""] + [f"- {t}" for t in sec["foreshadow"]]
        if _items(sec):
            L += ["", "**引用图 / 表：**", ""] + [f"- {_item_line(i, pending)}" for i in _items(sec)]
        if sec.get("code"):
            L += ["", "**关联代码（只作为你脑内映射）：**", ""] + [f"- {r}：{line}" for r, line in _code_roles(sec["code"])]
        return L
    if sec.get("paragraphs"):
        L += ["**写作要点**", ""]
        for p in sec["paragraphs"]:
            L.append(f"- **{_label(p)}**：{p.get('claim', '')}{_cite_tail(p)}")
            L += [f"  - {t}" for t in p.get("points", [])]
    if _items(sec):
        L += ["", "**图 / 表**", ""] + [f"- {_item_line(i, pending)}" for i in _items(sec)]
    if sec.get("methods"):
        L += ["", f"**{_methods_line(sec)}**"]
    if sec.get("boundaries"):
        L += ["", "**措辞边界**", ""] + [f"- {t}" for t in sec["boundaries"]]
    return L


def _check_records(o):
    """title_check / causal_design 记录（结论纪律），随大纲 md/docx 一起交付（评审 2026-10-14：只校验不渲染
    会让门禁逼出来的证据留在 yaml 里，写作技能读不到）。"""
    rows = []
    for sec in outline_check._nodes(o.get("results", [])):
        if sec.get("title_check") or sec.get("causal_design"):
            rows.append((sec["id"], str(sec.get("title_check") or "—"), str(sec.get("causal_design") or "—")))
    return rows


def to_md(o):
    pending = _pending(o)
    L = [f"# 写作大纲：{o.get('project', '')}", ""]
    for results, secs in ((False, o.get("methods", [])), (True, o.get("results", []))):
        for sec, depth, _ in _walk(secs):
            L += [f"{'#' * (depth + 2)} {sec['id']} {sec['title']}", ""]
            if results and sec.get("subtitle"):
                L += [f"（{sec['subtitle']}）", ""]
            body = _md_body(sec, results, pending)
            L += body + ([""] if body else [])
    rows = o.get("_coverage") or []
    if rows:
        L += ["## 图表引用清单", "", "| 资产 | panel | 首次引用 | 引用小节 | 状态 |", "|---|---|---|---|---|"]
        L += [f"| {a} | {p or '—'} | {f} | {'、'.join(s) or '—'} | {_status_zh(st)} |" for a, p, f, s, st in rows]
        L.append("")
    recs = _check_records(o)
    if recs:
        L += ["## 结论核对记录（title_check / causal_design）", "",
              "| 小节 | 全称主张的核对（title_check） | 因果表述的设计依据（causal_design） |", "|---|---|---|"]
        L += [f"| {i} | {tc} | {cd} |" for i, tc, cd in recs]
        L.append("")
    return "\n".join(L)


def _status_zh(st):
    return {"cited": "已引用", "pending": "已引用（外部绘制，待交付）", "NOT CITED": "未引用"}.get(
        st, st.replace("excluded: ", "不进稿件："))


def _pending(o):
    return {outline_check._asset(n, l) for r in o.get("pending", []) for n, l in (outline_check._expand(r) or [])}


# ------------------------------------------------------------------------------------------------ docx
def to_docx(o, path):
    import docx
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    d = docx.Document()
    sec0 = d.sections[0]
    sec0.page_height, sec0.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec0, side, Cm(2.54))
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    st.paragraph_format.space_after = Pt(2)

    def para(text="", level=0, bold=False, size=None):
        p = d.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.74 * level)
        p.paragraph_format.keep_with_next = bool(bold or text.endswith("："))
        for part in re.split(r"(`[^`]+`)", text):
            if not part:
                continue
            code = part.startswith("`") and part.endswith("`")
            r = p.add_run(part[1:-1] if code else part)
            r.bold = bold
            if size:
                r.font.size = Pt(size)
            if code:
                r.font.name = "Consolas"
                r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        return p

    pending = _pending(o)
    for results, secs in ((False, o.get("methods", [])), (True, o.get("results", []))):
        for sec, depth, _ in _walk(secs):
            h = para(f"{sec['id']} {sec['title']}", bold=True, size=HEAD_PT.get(depth, 10.5))
            h.paragraph_format.space_before = Pt(10 if depth == 0 else 6)
            if results and sec.get("subtitle"):
                subtitle = para(f"（{sec['subtitle']}）", size=11 if depth == 0 else 10.5)
                subtitle.paragraph_format.keep_with_next = True
            if not results:
                if sec.get("focus"):
                    para("写作重点：", bold=True)
                    for blk in sec["focus"]:
                        para(f"{blk['topic']}：", level=1)
                        for t in blk.get("points", []):
                            para(t, level=2)
                if sec.get("foreshadow"):
                    para("为后文铺垫：", level=1)
                    for t in sec["foreshadow"]:
                        para(t, level=2)
                if _items(sec):
                    para("引用图 / 表：", level=1)
                    for i in _items(sec):
                        para(_item_line(i, pending), level=2)
                if sec.get("code"):
                    para("关联代码（只作为你脑内映射）：", level=1)
                    for role, line in _code_roles(sec["code"]):
                        para(f"{role}：{line}", level=2)
                continue
            if sec.get("paragraphs"):
                para("写作要点", bold=True)
                for p in sec["paragraphs"]:
                    # one line per paragraph: bold label + claim + where it is shown; points follow as plain lines
                    q = d.add_paragraph()
                    q.paragraph_format.left_indent = Cm(0.74)
                    q.paragraph_format.keep_with_next = bool(p.get("points"))
                    q.add_run(_label(p) + "：").bold = True
                    q.add_run(p.get("claim", "") + _cite_tail(p))
                    for t in p.get("points", []):
                        para(t, level=2)
            if _items(sec):
                para("图 / 表", bold=True)
                for i in _items(sec):
                    para(_item_line(i, pending), level=1)
            if sec.get("methods"):
                # a one-line pointer, not a heading: bold but not kept with the next block
                para(_methods_line(sec), bold=True).paragraph_format.keep_with_next = False
            if sec.get("boundaries"):
                para("措辞边界", bold=True)
                for t in sec["boundaries"]:
                    para(t, level=1)
    rows = o.get("_coverage") or []
    if rows:
        h = para("图表引用清单", bold=True, size=14)
        h.paragraph_format.space_before = Pt(10)
        tab = d.add_table(rows=1, cols=5)
        tab.style = "Table Grid"
        tab.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for cell, t in zip(tab.rows[0].cells, ("资产", "panel", "首次引用", "引用小节", "状态")):
            cell.text = t
            cell.paragraphs[0].runs[0].bold = True
        for a, p, f, s, st in rows:
            for cell, t in zip(tab.add_row().cells, (a, p or "—", f, "、".join(s) or "—", _status_zh(st))):
                cell.text = t
    recs = _check_records(o)
    if recs:
        h = para("结论核对记录（title_check / causal_design）", bold=True, size=14)
        h.paragraph_format.space_before = Pt(10)
        tab = d.add_table(rows=1, cols=3)
        tab.style = "Table Grid"
        tab.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for cell, t in zip(tab.rows[0].cells, ("小节", "全称主张的核对（title_check）", "因果表述的设计依据（causal_design）")):
            cell.text = t
            cell.paragraphs[0].runs[0].bold = True
        for i, tc, cd in recs:
            for cell, t in zip(tab.add_row().cells, (i, tc, cd)):
                cell.text = t
    d.save(str(path))


# ------------------------------------------------------------------------------------------------ folders
def _caption_text(root, o, name):
    """Caption draft for figure ``name`` (fig2 / s05): its own file, or the '## Fig2' / '## S5' section."""
    head = f"Fig{int(name[3:])}" if name.startswith("fig") else f"S{int(name[1:])}"
    pat = re.compile(rf"^## {head}\s*$", re.M)
    for c in o.get("captions", []):
        p = root / c
        files = sorted(p.glob("*.md")) if p.is_dir() else [p] if p.is_file() else []
        for f in files:
            s = f.read_text(encoding="utf-8")
            m = pat.search(s)
            if m:
                nxt = re.compile(r"^## ", re.M).search(s, m.end())
                return s[m.start(): nxt.start() if nxt else len(s)].rstrip() + "\n"
    return None


def _copy_section(root, o, sec, results, dest):
    """Copy the outputs this section cites (figures, tables, captions); for Methods also the code chain the outline
    names (no automatic dependencies). Results carry no code: they point to the Methods sections instead."""
    figs, tabs = _refs(sec, results)
    pending = _pending(o)
    (dest / "figures").mkdir(parents=True, exist_ok=True)
    (dest / "tables").mkdir(exist_ok=True)
    (dest / "captions").mkdir(exist_ok=True)
    for name in figs:
        if name in pending:
            continue
        src = outline_check._source_json(root, o["figures_root"], name)
        for f in [src, *(src.with_name(f"{name}.{e}") for e in FIG_EXT)]:
            if f.is_file():
                shutil.copy2(f, dest / "figures" / f.name)
        cap = _caption_text(root, o, name)
        if cap:
            (dest / "captions" / f"{name}.md").write_text(cap, encoding="utf-8")
    for t in tabs:
        for e in TAB_EXT:
            f = root / o["tables_root"] / f"{t}.{e}"
            if f.is_file():
                shutil.copy2(f, dest / "tables" / f.name)
    if results or not sec.get("code"):
        return
    # only the files the outline names: flat by file name, files sharing a name keep their relative path;
    # 配置 / 回执 follow their copy mode (full / excerpt = named keys only / path = not copied)
    # A file may support several points; keep the union of its requested excerpt keys.
    grouped = {}
    for c in sec["code"]:
        f = outline_check.code_file(o, root, c.get("path"))
        if c.get("copy") == "path" or f is None:
            continue
        key = (f.resolve(), c.get("copy", "full"))
        if key not in grouped:
            grouped[key] = (dict(c, symbols=list(c.get("symbols") or [])), f)
        else:
            prior = grouped[key][0]
            prior["symbols"] = list(dict.fromkeys(prior["symbols"] + list(c.get("symbols") or [])))
    entries = list(grouped.values())
    names = [(_excerpt_name(f) if c.get("copy") == "excerpt" else f.name) for c, f in entries]
    written = set()
    for (c, f), name in zip(entries, names):
        relative = _archive_relative(c["path"], f, root, o)
        target = dest / "code" / (name if names.count(name) == 1 else relative.parent / name)
        code_root = (dest / "code").resolve()
        if not target.resolve().is_relative_to(code_root):
            raise ValueError(f"code destination escapes its section: {target}")
        if target in written:
            raise ValueError(f"conflicting code archive destinations: {target}")
        written.add(target)
        target.parent.mkdir(parents=True, exist_ok=True)
        if c.get("copy") == "excerpt":
            target.write_text(_excerpt(f, c.get("symbols") or [], c["path"]), encoding="utf-8")
        else:
            shutil.copy2(f, target)


def _archive_relative(path, f, root, o):
    """Keep listed relative paths on collisions, normalizing absolute / parent paths inside the archive."""
    rel = Path(path)
    if not rel.is_absolute() and ".." not in rel.parts:
        return rel
    source = f.resolve()
    for base in [Path(root), *(Path(root) / r for r in o.get("code_roots") or [])]:
        if source.is_relative_to(base.resolve()):
            return source.relative_to(base.resolve())
    # Explicit absolute sources outside the registered roots still keep their directory structure.
    return Path(*source.parts[1:])


def _excerpt_name(f):
    """report.json -> report.excerpt.json, report.yaml -> report.yaml.excerpt.json: the source suffix stays in the
    name so excerpts of report.json and report.yaml in one folder never collide."""
    if f.suffix == ".jsonl":
        return f"{f.stem}.excerpt.jsonl"
    return f"{f.stem}.excerpt.json" if f.suffix == ".json" else f"{f.name}.excerpt.json"


def _excerpt(f, keys, rel):
    """The named keys of a json / yaml file (dotted paths kept nested), or the jsonl records whose ``key`` is named,
    with the source path on top so the excerpt can be traced back. Loading and key paths are outline_check's, so a
    file that passed the check is read the same way here."""
    import json
    data = outline_check.load_structured(f)
    if f.suffix == ".jsonl":
        rows = [r for r in data if r.get("key") in set(keys)]
        head = json.dumps({"_excerpt_of": rel, "_keys": list(keys)}, ensure_ascii=False)
        return "\n".join([head] + [json.dumps(r, ensure_ascii=False) for r in rows]) + "\n"
    out = {"_excerpt_of": rel}
    for k in keys:
        found, value = outline_check.dotted(data, k)
        if not found:
            continue
        dst = out
        *heads, last = str(k).split(".")
        for part in heads:
            dst = dst.setdefault(part, {})
        dst[last] = value
    return json.dumps(out, ensure_ascii=False, indent=2, default=str) + "\n"


def build(outline_path, project, out, figkit_toml=None):
    root, out = Path(project).resolve(), Path(out).resolve()
    if out == root or root in out.parents:
        raise ValueError(f"output {out} must be outside the project {root}")
    issues = outline_check.check(outline_path, root, figkit_toml)
    if issues:
        raise RuntimeError("outline_check reported issues:\n" + "\n".join(issues))
    o = yaml.safe_load(Path(outline_path).read_text(encoding="utf-8"))
    o["_coverage"] = outline_check.coverage(o, root)
    out.mkdir(parents=True, exist_ok=True)
    md = out / "写作大纲.md"
    md.write_text(to_md(o), encoding="utf-8")
    dx = out / "写作大纲.docx"
    to_docx(o, dx)
    for results, top, secs in ((False, "2_Methods", o.get("methods", [])), (True, "3_Results", o.get("results", []))):
        for sec, _, parents in outline_check._tree(secs):
            # subsections live inside their parent's folder: 3_Results/3.1_x/3.1.2_y/
            dest = out / top
            for par in parents:
                dest = dest / f"{par['id']}_{par['short']}"
            dest = dest / f"{sec['id']}_{sec['short']}"
            if _refs(sec, results) != ([], []) or sec.get("code"):
                _copy_section(root, o, sec, results, dest)
            else:
                dest.mkdir(parents=True, exist_ok=True)
    return dict(docx=str(dx), md=str(md), out=str(out))


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("outline")
    ap.add_argument("--project", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--figkit-toml")
    a = ap.parse_args(argv)
    res = build(a.outline, a.project, a.out, a.figkit_toml)
    print(res["docx"]); print(res["md"]); print(res["out"])
    return 0


if __name__ == "__main__":
    sys.exit(main())
