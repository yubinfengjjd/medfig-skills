"""End-to-end assembly checks use synthetic fixtures, never the real project."""
from pathlib import Path
import sys

import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import build_manuscript as build
import check_finalize as check
import test_finalize as legacy_tests


@pytest.fixture
def assembled(tmp_path):
    legacy_tests.Assembly().fixture(tmp_path)
    path = tmp_path / '02_figures/out/tables/ST01.md'
    path.write_text('| A | B |\n|---|---|\n| first | 1 |\n\n'
                    '| C | D | E |\n|---|---|---|\n| second | 2 | 3 |\n\n'
                    '| F | G |\n|---|---|\n| third | 4 |\n', encoding='utf-8')
    output = tmp_path / 'output'
    build.build(tmp_path, output_dir=output, update_tracking=False)
    return tmp_path, output, output / f'{tmp_path.name}_Supplementary_v1.docx'


def test_checker_accepts_all_three_blocks(assembled):
    root, output, _ = assembled
    result = check.check_assembly(root, output)
    assert not result['errors'], result['errors']
    assert any('3 blocks' in entry for entry in result['information'])
    assert [x['block'] for x in result['documents']['supp']['table_layouts']] == [1, 2, 3]


@pytest.mark.parametrize('mutation', ['cell', 'width', 'font', 'missing', 'duplicate', 'gap', 'note'])
def test_checker_rejects_later_block_corruption(assembled, mutation):
    root, output, path = assembled
    doc = Document(path)
    second = doc.tables[1]
    if mutation == 'cell':
        second.cell(1, 0).paragraphs[0].runs[0].text = 'corrupted'
    elif mutation == 'width':
        second.columns[0].width = build.Cm(1)
    elif mutation == 'font':
        second.cell(1, 0).paragraphs[0].runs[0].font.size = build.Pt(12)
    elif mutation == 'missing':
        second._tbl.getparent().remove(second._tbl)
    elif mutation == 'duplicate':
        from copy import deepcopy
        second._tbl.addnext(deepcopy(second._tbl))
    elif mutation == 'gap':
        gap = second._tbl.getprevious()
        gap.getparent().remove(gap)
    else:
        doc.paragraphs[-2].runs[0].text = 'Incorrect notes.'
    doc.save(path)
    assert check.check_assembly(root, output)['errors']
