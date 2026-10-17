"""Publication IDs and native Word tables; all fixtures are synthetic."""
import json
from pathlib import Path
import sys

import pytest
from docx.oxml.ns import qn

SCRIPTS = Path(__file__).resolve().parents[1] / 'scripts'
sys.path.insert(0, str(SCRIPTS))
import build_manuscript as build
import test_finalize as legacy_tests

TABLE = '| Item | Value |\n| --- | --- |\n| A | 0.962 |\n'


@pytest.fixture
def project(tmp_path):
    legacy_tests.Assembly().fixture(tmp_path)
    return tmp_path


def mapped_table(project, title='Supplementary Table S1', registered=None, name='ST10'):
    folder = project / '02_figures/out/tables'
    (folder / 'ST01.md').unlink(missing_ok=True)
    path = folder / (name + '.md')
    path.write_text('  **' + title + '.** Title.\n\n' + TABLE, encoding='utf-8')
    if registered is not None:
        path.with_suffix('.source.json').write_text(json.dumps({
            'presentation_trim': {'display_number': registered}}), encoding='utf-8')
    return path


def test_st10_maps_to_display_s9(project):
    path = mapped_table(project, 'Supplementary Table S9', 'Supplementary Table S9')
    notes = project / '04_manuscript/Legends/Table_Titles_Notes_EN.md'
    notes.write_text(notes.read_text(encoding='utf-8') + '\n' + '\n'.join(
        f'**Supplementary Table S{i}. Extra.** Note.' for i in range(2, 10)), encoding='utf-8')
    folder = path.parent
    for i in range(1, 9):
        (folder / f'ST{i:02}.md').write_text(TABLE, encoding='utf-8')
    assert build.project_model(project)['tables']['S9'] == path
    assert 'S10' not in build.project_model(project)['tables']


@pytest.mark.parametrize('registered', ['Supplementary Table S2', 'Table S1'])
def test_display_provenance_mismatch_fails(project, registered):
    mapped_table(project, registered=registered)
    with pytest.raises(ValueError, match='display/provenance'):
        build.project_model(project)


def test_duplicate_display_number_fails(project):
    mapped_table(project)
    (project / '02_figures/out/tables/ST11.md').write_text(
        '**Supplementary Table S1.** Duplicate.\n' + TABLE, encoding='utf-8')
    with pytest.raises(ValueError, match='Duplicate.*S1'):
        build.project_model(project)


def test_duplicate_fallback_number_fails(project):
    (project / '02_figures/out/tables/ST1.md').write_text(TABLE, encoding='utf-8')
    with pytest.raises(ValueError, match='Duplicate.*S1'):
        build.project_model(project)


def test_no_title_keeps_filename_fallback(project):
    assert set(build.project_model(project)['tables']) == {'1', '2', 'S1'}


def test_main_title_and_arbitrary_filename(project):
    folder = project / '02_figures/out/tables'
    (folder / 'T1.md').unlink()
    path = folder / 'stable-source.md'
    path.write_text('  **Table 1.** Main.\n' + TABLE, encoding='utf-8')
    path.with_suffix('.source.json').write_text(json.dumps({
        'presentation_trim': {'display_number': 'Table 1'}}), encoding='utf-8')
    assert build.project_model(project)['tables']['1'] == path


def test_empty_provenance_number_is_optional(project):
    mapped_table(project, registered='')
    assert 'S1' in build.project_model(project)['tables']


def render(tmp_path, text, module=build):
    path = tmp_path / 'table.md'
    path.write_text(text, encoding='utf-8')
    model = {'tables': {'S5': path}, 'table_captions': {'S5': 'Supplementary Table S5. Data. Notes.'}}
    doc = module.new_document()
    warnings, information, layouts = [], [], []
    result = module.add_table(doc, 'S5', model, warnings, information, layouts)
    return doc, warnings, information, layouts, result


@pytest.mark.parametrize('count', [2, 3])
def test_multiple_blocks_have_one_caption_and_one_gap(tmp_path, count):
    doc, warnings, information, layouts, result = render(tmp_path, '\n'.join([TABLE] * count))
    nodes = list(doc.element.body)[:-1]
    assert [n.tag for n in nodes] == [qn('w:p')] + [
        qn('w:tbl') if i % 2 == 0 else qn('w:p') for i in range(2 * count - 1)] + [qn('w:p')]
    assert [p.text for p in doc.paragraphs] == ['Supplementary Table S5. Data.'] + [''] * (count - 1) + [' Notes.']
    assert len(doc.tables) == count
    assert [layout['block'] for layout in layouts] == list(range(1, count + 1))
    assert any(f'{count} blocks' in entry for entry in information)
    assert len(result) == count
    assert not warnings


def test_each_block_gets_own_layout_and_numbered_warning(tmp_path):
    wide = '| ' + ' | '.join(['LongHeader' * 10] * 5) + ' |\n'
    wide += '| ' + ' | '.join(['---'] * 5) + ' |\n'
    wide += '| ' + ' | '.join(['LongWord' * 20] * 5) + ' |\n'
    _, warnings, information, layouts, _ = render(tmp_path, TABLE + '\n' + wide + '\n' + wide)
    assert [layout['font_size_pt'] for layout in layouts] == [8, 6.5, 6.5]
    assert all(f'Table S5 block {i}' in warnings[i - 2] for i in (2, 3))
    assert all(any(f'Table S5 block {i}' in entry for entry in information) for i in (1, 2, 3))


@pytest.mark.parametrize('bad', [
    '| Item | Value |\n|---|---|\n',
    '| Item | Value |\n|---|---|\n| bad |\n',
    '| Item | Value |\n|---|\n| A | 1 |\n',
    '| Item | Value |\n|oops|---|\n| A | 1 |\n',
    '| Item | Value |\n',
    '| | |\n|---|---|\n| | |\n',
])
def test_bad_or_empty_second_block_fails_before_rendering(tmp_path, bad):
    path = tmp_path / 'table.md'
    path.write_text(TABLE + '\n' + bad, encoding='utf-8')
    doc = build.new_document()
    with pytest.raises(ValueError, match=r'Table S5 block 2'):
        build.add_table(doc, 'S5', {'tables': {'S5': path}, 'table_captions': {'S5': 'Table S5. Data.'}})
    assert not doc.tables


def test_no_table_fails(tmp_path):
    with pytest.raises(ValueError, match='Table S5'):
        render(tmp_path, '**Supplementary Table S5.** Empty.')


@pytest.mark.parametrize('text', [TABLE,
    '| Label | Math | Empty |\n|---|---|---|\n| a\\|b | $x_1$ | |\n',
    '| ' + 'LongHeader' * 20 + ' | Value |\n|---|---|\n| ' + 'LongWord' * 40 + ' | 1 |\n'])
def test_single_block_bytes_match_before_t3c(tmp_path, text):
    namespace = dict(vars(build))
    baseline = Path(__file__).with_name('legacy_add_table.txt').read_text(encoding='utf-8')
    exec(compile(baseline, 'legacy_add_table.txt', 'exec'), namespace)
    from types import SimpleNamespace
    legacy = SimpleNamespace(**namespace)
    current = render(tmp_path, text)
    before = render(tmp_path, text, legacy)
    assert build.stable_save(current[0], None) == build.stable_save(before[0], None)
    assert current[1:] == before[1:]
