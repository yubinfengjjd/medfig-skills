"""Offline regression suite. Fixtures use the default system temporary directory."""
from pathlib import Path
import json
import subprocess
import sys
import tempfile
import unittest
import hashlib
import shutil
from PIL import Image
from docx import Document
from docx.oxml.ns import qn

import check_finalize as check
import build_manuscript as build


NL = chr(10)


class Invariants(unittest.TestCase):
    def test_no_skill_local_temp_directory(self):
        with tempfile.TemporaryDirectory() as d:
            scripts = Path(d) / 'scripts'
            scripts.mkdir()
            for name in ('build_manuscript.py', 'check_finalize.py', 'test_finalize.py'):
                shutil.copyfile(Path(__file__).with_name(name), scripts / name)
            result = subprocess.run([sys.executable, '-B', '-X', 'utf8',
                                     str(scripts / 'test_finalize.py'),
                                     'Invariants.test_section_cannot_hide_number_swap'],
                                    capture_output=True, text=True, encoding='utf-8')
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            self.assertEqual({p.name for p in Path(d).iterdir()}, {'scripts'})
            self.assertEqual({p.name for p in scripts.iterdir()},
                             {'build_manuscript.py', 'check_finalize.py', 'test_finalize.py'})

    def test_numeric_multiset_not_set(self):
        self.assertTrue(check.compare_text('0.5 0.5 2', '0.5 2 2')['errors'])
        self.assertTrue(check.compare_text('−0.03 +0.04', '0.03 +0.04')['errors'])
        self.assertTrue(check.compare_text('1e-3', '1e-4')['errors'])

    def test_qualifier_removal(self):
        before = 'A single evaluated seed was used; evidence was not conclusive. The 95% CI crossed zero.'
        self.assertFalse(check.compare_text(before, before)['errors'])
        self.assertTrue(check.compare_text(before, before.replace('single evaluated seed', 'seed'))['errors'])
        self.assertTrue(check.compare_text(before, before.replace('not conclusive', 'conclusive'))['errors'])
        self.assertTrue(check.compare_text(before, before.replace('crossed zero', 'excluded zero'))['errors'])

    def test_units_citations_and_figure_ids(self):
        for a, b in [('5 mm', '5 cm'), ('(PMID: 12345678)', '(DOI: 12345678)'),
                     ('(Fig. S1a)', '(Fig. S1b)'), ('P < 0.05', 'P > 0.05'),
                     ('5 mg', '5 g'), ('5 mm and 3 cm', '5 cm and 3 mm'),
                     ('The 95% CI included zero.', 'The 95% CI excluded zero.')]:
            self.assertTrue(check.compare_text(a, b)['errors'], (a, b))

    def test_word_band(self):
        text = ' '.join(['word'] * 100)
        self.assertFalse(check.compare_text(text, ' '.join(['word'] * 90))['errors'])
        self.assertTrue(check.compare_text(text, ' '.join(['word'] * 89))['errors'])
        self.assertTrue(check.compare_text(text, ' '.join(['word'] * 111))['errors'])

    def test_word_count_excludes_standalone_panel_labels(self):
        # 2 title tokens + 159 body tokens; three panel labels must not add words.
        caption = '**Figure 1.** (a) ' + 'word ' * 53 + '(b) ' + 'word ' * 53 + '(c) ' + 'word ' * 53
        self.assertEqual(check.word_count(caption), 161)
        for labels in ('(a) (a-c) (a–c)', '（a） （a-c） （a–c）',
                       '(a), (b); (c): (a–c).',
                       '(a)，（b）；(c)：（a–c）。'):
            with self.subTest(labels=labels):
                self.assertEqual(check.word_count(labels + ' Observed response.'), 2)
        # Scientific parentheticals and embedded function notation remain content.
        self.assertEqual(check.word_count('Observed (n=12) (control) response x(a).'), 5)

    def test_section_cannot_hide_number_swap(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d)
            for sub in ('before', 'after'):
                (p / sub).mkdir()
            (p/'before/A_EN.md').write_text('Value 1.', encoding='utf-8')
            (p/'before/B_EN.md').write_text('Value 2.', encoding='utf-8')
            (p/'after/A_EN.md').write_text('Value 2.', encoding='utf-8')
            (p/'after/B_EN.md').write_text('Value 1.', encoding='utf-8')
            self.assertTrue(check.compare_paths(p/'before', p/'after')['errors'])

    def test_references_order_and_dedup(self):
        text = '(PMID: 22222222) then (DOI: 10.1/test); (PMID: 11111111; PMID: 22222222)'
        self.assertEqual(build.citation_ids(text), ['PMID: 22222222', 'DOI: 10.1/test', 'PMID: 11111111'])

    def test_captions_duplicate_missing(self):
        ok = '**Figure 1. A.** Body.\n**Figure 2. B.** Body.'
        self.assertEqual(set(build.captions(ok, 'figure')), {'1', '2'})
        with self.assertRaises(ValueError):
            build.captions(ok + '\n**Figure 1. Duplicate.** Body.', 'figure')
        self.assertTrue(check.match_ids(['1', '2'], ['1'], 'figure'))

    def test_cli_rejects_injected_changes(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d)
            a = p/'before.md'
            b = p/'after.md'
            a.write_text('The single-run BACC was 0.962 and was not conclusive.', encoding='utf-8')
            for bad in ['The single-run BACC was 0.963 and was not conclusive.',
                        'The BACC was 0.962 and was not conclusive.']:
                b.write_text(bad, encoding='utf-8')
                r = subprocess.run([sys.executable, '-X', 'utf8', str(Path(check.__file__)),
                                    '--before', str(a), '--after', str(b)], capture_output=True, text=True, encoding='utf-8')
                self.assertEqual(r.returncode, 1, r.stdout + r.stderr)
                self.assertTrue(json.loads(r.stdout)['errors'])

    def test_apply_archives_english_and_preserves_chinese(self):
        with tempfile.TemporaryDirectory() as d:
            root = Path(d); original=root/'04_manuscript'; staged=root/'staged'
            original.mkdir(); staged.mkdir()
            old = 'Unfortunately, a single-run value was 0.962 and was not conclusive.'
            new = 'The single-run value was 0.962 and was not conclusive.'
            (original/'Sample_EN.md').write_text(old, encoding='utf-8')
            (staged/'Sample_EN.md').write_text(new, encoding='utf-8')
            (original/'Sample_ZH.md').write_bytes('中文不动'.encode('utf-8'))
            changes = [{'file':'Sample_EN.md', **x, 'reason':'Remove emotional preface'} for x in check.paragraph_changes(old,new)]
            (root/'changes.json').write_text(json.dumps(changes), encoding='utf-8')
            (root/'decisions.md').write_text('新增 0 项；保持原文证据限定。', encoding='utf-8')
            result=check.apply_polish(original,staged,root/'changes.json',root/'decisions.md',root,'2026-10-15',['single-run','not conclusive'])
            self.assertFalse(result['errors'], result)
            self.assertEqual((root/'_archive/2026-10-15/04_manuscript/Sample_EN.md').read_text(encoding='utf-8'),old)
            self.assertEqual((original/'Sample_EN.md').read_text(encoding='utf-8'),new)
            self.assertEqual((original/'Sample_ZH.md').read_bytes(),'中文不动'.encode('utf-8'))

    def test_incomplete_or_unreasoned_log_rejected(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d); a=p/'a.md'; b=p/'b.md'; changes=p/'changes.json'
            a.write_text('We describe a clear finding.',encoding='utf-8'); b.write_text('We report a clear finding.',encoding='utf-8')
            for data in [[], [{'file':'a.md', **check.paragraph_changes(a.read_text(),b.read_text())[0]}]]:
                changes.write_text(json.dumps(data),encoding='utf-8')
                self.assertTrue(check.compare_paths(a,b,changes=changes)['errors'])


class Assembly(unittest.TestCase):
    def test_batch10_section_end_mixed_assets_and_plain_headings(self):
        with tempfile.TemporaryDirectory() as d:
            root = Path(d); self.fixture(root); out = root/'output'
            (root/'04_manuscript/Methods/Methods_EN.md').write_text(
                '# Methods\n\n## 2.3 Parent\n\nFirst Table 2 then Fig. 1.\n\nParent tail.'
                '\n\n### 2.3.1 Child\n\nChild cites Table 1.\n\nChild tail.'
                '\n\n### 2.3.2 Sibling\n\nSibling tail.\n\n## 2.4 Next\n\nNext body.', encoding='utf-8')
            manifest = build.build(root, output_dir=out, update_tracking=False)
            doc = Document(out/f'{root.name}_Manuscript_v1.docx')
            texts = [p.text for p in doc.paragraphs]
            t1 = next(i for i,t in enumerate(texts) if t.startswith('Table 1.'))
            t2 = next(i for i,t in enumerate(texts) if t.startswith('Table 2.'))
            fig = next(i for i,t in enumerate(texts) if t.startswith('Figure 1.'))
            self.assertLess(texts.index('Child tail.'), t1)
            self.assertLess(t1, texts.index('2.3.2 Sibling'))
            self.assertLess(texts.index('Sibling tail.'), t2)
            self.assertLess(t2, fig); self.assertLess(fig, texts.index('2.4 Next'))
            self.assertNotIn('References', texts)
            self.assertFalse(doc.element.xpath('.//w:pStyle[starts-with(@w:val,"Heading") or starts-with(@w:val,"Title")]|.//w:outlineLvl'))
            for text, size in [('[Manuscript title]',16), ('Abstract',15), ('Keywords',15), ('2.3 Parent',12)]:
                p = next(p for p in doc.paragraphs if p.text == text)
                self.assertEqual(p.style.name, 'Normal')
                self.assertTrue(all(r.bold and r.font.size.pt == size for r in p.runs))
            for label in ['Abstract','Keywords','Highlights',*build.SECTIONS]:
                i = texts.index(label); self.assertEqual(texts[i-1], '')
            self.assertFalse(any('bibliographic' in w.lower() or 'unverified' in w.lower() for w in manifest['warnings']))
            self.assertFalse(check.check_assembly(root,out)['errors'])

    def test_batch10_checker_rejects_style_outline_blank_and_references(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d); self.fixture(root); out=root/'output'
            build.build(root,output_dir=out,update_tracking=False)
            path=out/f'{root.name}_Manuscript_v1.docx'; original=path.read_bytes()
            from docx.oxml import OxmlElement
            for mutation, phrase in [('style','Heading/Title'), ('outline','outline'),
                                     ('blank','blank'), ('refs','References'), ('size','heading')]:
                path.write_bytes(original); doc=Document(path)
                p=next(p for p in doc.paragraphs if p.text=='Introduction')
                if mutation=='style': p.style='Heading 1'
                if mutation=='outline': p._p.get_or_add_pPr().append(OxmlElement('w:outlineLvl'))
                if mutation=='blank': p._p.getparent().remove(p._p.getprevious())
                if mutation=='refs': doc.add_paragraph('References')
                if mutation=='size': p.runs[0].font.size=build.Pt(12)
                doc.save(path)
                errors=check.check_document(path,build.project_model(root),'main')['errors']
                self.assertTrue(any(phrase.lower() in e.lower() for e in errors),(mutation,errors))

    def test_batch9_page_and_typography(self):
        doc = build.new_document()
        self.assertEqual(len(doc.sections), 1)
        sec = doc.sections[0]
        self.assertAlmostEqual(sec.page_width.cm, 21, delta=.01)
        self.assertAlmostEqual(sec.page_height.cm, 29.7, delta=.01)
        for v in (sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin):
            self.assertAlmostEqual(v.cm, 3.2, delta=.01)
        for name, size in [('Normal', 12), ('Title', 16), ('Heading 1', 15), ('Heading 2', 12)]:
            self.assertEqual(doc.styles[name].font.name, 'Arial')
            self.assertEqual(doc.styles[name].font.size.pt, size)
            self.assertEqual(doc.styles[name].paragraph_format.line_spacing, 1)
        self.assertNotIn('lnNumType', doc.element.xml)
        self.assertFalse(sec._sectPr.xpath('./w:headerReference|./w:footerReference'))

    def test_batch9_full_width_natural_figure_and_caption_runs(self):
        with tempfile.TemporaryDirectory() as d:
            root = Path(d); image = root/'figure.png'
            Image.new('RGB', (1200, 2000), 'white').save(image)
            for caption in ('Figure 1. Title sentence. (a) Body. Tail.',
                            'Figure 1. Title sentence. ' + 'Body ' * 1000):
                doc = build.new_document(); warnings = []
                build.add_figure(doc, '1', image, caption, root, warnings)
                shape = doc.inline_shapes[0]
                self.assertAlmostEqual(shape.width.cm, 14.6, delta=.001)
                self.assertAlmostEqual(shape.height/shape.width, 2000/1200, delta=.001)
                picture, legend = doc.paragraphs[-2:]
                self.assertFalse(picture.paragraph_format.keep_with_next)
                self.assertFalse(picture.paragraph_format.page_break_before)
                self.assertFalse(legend.paragraph_format.page_break_before)
                self.assertFalse(legend.paragraph_format.keep_together)
                self.assertEqual(legend.text, caption)
                self.assertEqual(''.join(r.text for r in legend.runs if r.bold),
                                 'Figure 1. Title sentence.')
                self.assertTrue(all(r.bold is False for r in legend.runs[1:]))
                self.assertFalse(warnings)

    def test_batch10_single_table_floor_wrap_and_preserved_data(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d); path=root/'table.md'
            headers=['ID']+[f'Column{i}' for i in range(1,9)]
            rows=[['A']+['MMMMMMMMMMMM']*8, ['B']+['123.456789']*8]
            path.write_text('|'+'|'.join(headers)+'|\n|'+'---|'*len(headers)+'\n'+
                            '\n'.join('|'+'|'.join(r)+'|' for r in rows),encoding='utf-8')
            original=path.read_bytes()
            model={'tables':{'1':path},'table_captions':{'1':'Table 1. Title sentence. Notes remain regular.'}}
            doc=build.new_document(); warnings=[]; info=[]; layouts=[]
            build.add_table(doc,'1',model,warnings=warnings,information=info,layouts=layouts)
            self.assertEqual(len(doc.tables),1)
            table=doc.tables[0]
            self.assertEqual([[c.text for c in r.cells] for r in table.rows],[headers,*rows])
            self.assertAlmostEqual(sum(c.width.cm for c in table.columns),14.6,delta=.01)
            self.assertEqual({r.font.size.pt for row in table.rows for c in row.cells
                             for p in c.paragraphs for r in p.runs},{6.5})
            self.assertTrue(layouts[0]['word_wrap_columns']); self.assertTrue(warnings)
            self.assertNotIn('(continued)',doc.element.xml)
            self.assertFalse(table._tbl.xpath('.//w:cantSplit'))
            self.assertTrue(table.rows[0]._tr.xpath('./w:trPr/w:tblHeader'))
            self.assertEqual(doc.paragraphs[0].text,'Table 1. Title sentence.')
            self.assertEqual(doc.paragraphs[-1].text,' Notes remain regular.')
            self.assertFalse(any(r.bold for r in doc.paragraphs[-1].runs))
            self.assertEqual(path.read_bytes(),original)

    def test_batch10_font_steps_and_long_word_floor(self):
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'table.md'
            path.write_text('| ID | Value |\n|---|---|\n| A | '+'9'*400+' |',encoding='utf-8')
            model={'tables':{'1':path},'table_captions':{'1':'Table 1. Data.'}}
            doc=build.new_document(); warnings=[]
            layout=build.add_table(doc,'1',model,warnings=warnings)
            self.assertEqual(layout['font_size_pt'],6.5)
            self.assertEqual(doc.tables[0].cell(1,1).text,'9'*400)
            self.assertTrue(warnings)
            self.assertEqual(layout['word_wrap_columns'][0]['column'],2)
            for size in (8,7.5,7,6.5):
                cols=[['ID','A']]+[[f'C{i}','MMMMMMMMMM'] for i in range(7)]
                required=sum(build.column_width_requirements(cols,size)[0])/20
                previous=sum(build.column_width_requirements(cols,size+.5)[0])/20
                plan=build.table_plan(cols,(required+previous)/2)
                self.assertEqual(plan['size'],size)

    def test_batch9_first_citation_order_abstract_and_supp_exclusion(self):
        with tempfile.TemporaryDirectory() as d:
            root = Path(d); self.fixture(root)
            abstract = root/'04_manuscript/Abstract/Abstract_EN.md'
            abstract.write_text('## Abstract\n\nFirst mention (Fig. 1; Table 2).'
                                '\n\n## Keywords\n\nImaging\n\n## Highlights\n\n- Finding.', encoding='utf-8')
            intro = root/'04_manuscript/Introduction/Introduction_EN.md'
            intro.write_text('## Introduction\n\nSupplement only (Supplementary Figure S1; Table S1).'
                             '\n\nCohort (Table 1).', encoding='utf-8')
            out = root/'output'; build.build(root, output_dir=out, update_tracking=False)
            doc = Document(out/f'{root.name}_Manuscript_v1.docx')
            body = list(doc.element.body)
            cite = next(p for p in doc.paragraphs if p.text == 'First mention (Fig. 1; Table 2).')
            i = body.index(cite._p)
            self.assertEqual(body[i+1].tag, qn('w:p'))
            self.assertTrue(body[i+2].xpath('.//w:drawing'))
            self.assertEqual(body[i+3].xpath('.//w:t')[0].text, 'Figure 1. Main.')
            self.assertTrue('Table 2.' in ''.join(body[i+5].itertext()))
            self.assertEqual(body[i+6].tag, qn('w:tbl'))
            self.assertNotIn('References', [p.text for p in doc.paragraphs])
            self.assertFalse(check.check_assembly(root, out)['errors'])

    def test_batch9_checker_rejects_layout_and_caption_corruption(self):
        with tempfile.TemporaryDirectory() as d:
            root = Path(d); self.fixture(root); out = root/'output'
            build.build(root, output_dir=out, update_tracking=False)
            path = out/f'{root.name}_Manuscript_v1.docx'; original = path.read_bytes()
            from docx.oxml import OxmlElement
            for mutation, expected in [('width', 'width'), ('bold', 'bold'),
                                       ('numbering', 'line numbering'), ('orientation', 'portrait'),
                                       ('placement', 'first citation'), ('page', 'PAGE')]:
                path.write_bytes(original); doc = Document(path)
                if mutation == 'width': doc.inline_shapes[0].width = build.Cm(8)
                if mutation == 'bold':
                    next(p for p in doc.paragraphs if p.text.startswith('Figure 1.')).runs[-1].bold = True
                if mutation == 'numbering': doc.sections[0]._sectPr.append(OxmlElement('w:lnNumType'))
                if mutation == 'orientation': doc.sections[0]._sectPr.pgSz.set(qn('w:orient'), 'landscape')
                if mutation == 'placement':
                    picture = doc.inline_shapes[0]._inline.getparent().getparent().getparent()
                    doc.element.body.insert(len(doc.element.body)-1, picture)
                if mutation == 'page':
                    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
                    doc.sections[0].footer.paragraphs[0]._p.append(fld)
                doc.save(path)
                errors = check.check_document(path, build.project_model(root), 'main')['errors']
                self.assertTrue(any(expected in e for e in errors), (mutation, errors))

    def test_batch9_compact_front_and_no_breaks(self):
        with tempfile.TemporaryDirectory() as d:
            root = Path(d); self.fixture(root); out = root/'output'
            build.build(root, output_dir=out, update_tracking=False)
            doc = Document(out/f'{root.name}_Manuscript_v1.docx')
            texts = [p.text for p in doc.paragraphs]; index = texts.index('Abstract')
            self.assertEqual(index, 5)
            self.assertLessEqual(sum(not t for t in texts[:index]), 3)
            self.assertEqual(doc.paragraphs[0].runs[0].font.size.pt, 16)
            self.assertEqual(doc.paragraphs[1].runs[0].font.size.pt, 10)
            self.assertTrue(doc.paragraphs[1].runs[0].bold)
            self.assertFalse(doc.element.xpath('.//w:br[@w:type="page"]'))
            self.assertFalse(doc.element.xpath('.//w:pageBreakBefore[@w:val="1"]'))

    def test_batch9_headings_math_and_theme_resolve_to_arial(self):
        doc = build.new_document()
        build.add_markdown(doc, '## Subsection\n\nA body paragraph.')
        heading = doc.paragraphs[0]
        self.assertEqual(heading.style.font.size.pt, 12)
        self.assertTrue(all(r.bold for r in heading.runs))
        self.assertFalse(any('theme' in a.lower()
                             for a in heading.style.element.get_or_add_rPr().rFonts.attrib))
        build.add_caption(doc, r'Figure 1. Model \(p_c\). (a) Value \(x\).')
        p = doc.paragraphs[-1]
        math_runs = p._p.xpath('.//m:r')
        self.assertTrue(math_runs)
        for run in math_runs:
            self.assertEqual(run.find('./'+qn('w:rPr')+'/'+qn('w:rFonts')).get(qn('w:ascii')), 'Arial')
        self.assertEqual(math_runs[0].find('./'+qn('w:rPr')+'/'+qn('w:b')).get(qn('w:val')), '1')
        self.assertEqual(math_runs[-1].find('./'+qn('w:rPr')+'/'+qn('w:b')).get(qn('w:val')), '0')

    def test_batch10_checker_single_table_and_corruption(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d); self.fixture(root); out=root/'output'
            source=root/'02_figures/out/tables/ST01.md'
            source.write_text('|ID|'+'|'.join('C'+str(i) for i in range(8))+'|\n|'+
                              '---|'*9+'\n|A|'+'MMMMMMMMMMMM|'*8,encoding='utf-8')
            manifest=build.build(root,output_dir=out,update_tracking=False)
            path=out/f'{root.name}_Supplementary_v1.docx'; original=path.read_bytes()
            self.assertEqual(len(Document(path).tables),1)
            result=check.check_assembly(root,out)
            self.assertFalse(result['errors'],result)
            self.assertEqual(manifest['information'],result['information'])
            process=next((out/'_process').iterdir())
            self.assertIn('6.5 pt',build.read(process/'表格断行清单.md'))
            from copy import deepcopy
            from docx.oxml import OxmlElement
            for change,phrase in [('cell','contents'),('duplicate','single'),('caption','caption'),
                                  ('paper','A4'),('width','width'),('font','font'),('row','cross-page')]:
                path.write_bytes(original); doc=Document(path)
                if change=='cell': doc.tables[0].cell(1,0).text='Changed ID'
                if change=='duplicate':
                    table=doc.tables[0]._tbl
                    table.addnext(deepcopy(table)); table.addnext(deepcopy(table.getprevious()))
                if change=='caption':
                    next(p for p in doc.paragraphs if p.text.startswith('Supplementary Table')).text += ' (continued)'
                if change=='paper': doc.sections[0].page_height=build.Cm(27.9)
                if change=='width': doc.tables[0].columns[0].width=build.Pt(1)
                if change=='font': doc.tables[0].cell(1,0).paragraphs[0].runs[0].font.size=build.Pt(6)
                if change=='row': doc.tables[0].rows[1]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                doc.save(path)
                errors=check.check_document(path,build.project_model(root),'supp')['errors']
                self.assertTrue(any(phrase in e for e in errors),(change,errors))

    def test_v03_cell_escapes_decode_only_outside_math(self):
        with tempfile.TemporaryDirectory() as d:
            root = Path(d); self.fixture(root)
            path = root/'02_figures/out/tables/T1.md'
            path.write_text(r'| Metric\_label | Value |' + '\n|---|---|\n'
                            + r'| \|Coverage − nominal\| \_x\_ \*x\* \\ | $p_c$ and \(\alpha\) |',
                            encoding='utf-8')
            source = path.read_bytes(); out = root/'output'
            build.build(root, output_dir=out, update_tracking=False)
            doc = Document(out/f'{root.name}_Manuscript_v1.docx')
            self.assertEqual(doc.tables[0].cell(0, 0).text, 'Metric_label')
            self.assertEqual(doc.tables[0].cell(1, 0).text, '|Coverage − nominal| _x_ *x* \\')
            math = doc.tables[0].cell(1, 1)._tc.xpath('.//m:oMath')
            self.assertEqual(len(math), 2)
            from lxml import etree
            for actual, expr in zip(math, ('p_c', r'\alpha')):
                # Font size is a layout property, not a change to the OMML expression.
                for prop in list(actual.iter(qn('w:rPr'))): prop.getparent().remove(prop)
                self.assertEqual(etree.tostring(actual, method='c14n', exclusive=True),
                                 etree.tostring(build.math_element(expr), method='c14n', exclusive=True))
            self.assertEqual(path.read_bytes(), source)
            self.assertFalse(check.check_assembly(root, out)['errors'])
            # Escaped literal formatting is not interpreted as bold; math still fails closed.
            p = doc.add_paragraph()
            build.add_inline(p, r'\*\*literal\*\* **bold**', decode_escapes=True)
            self.assertEqual(p.text, '**literal** bold')
            self.assertFalse(any(r.bold for r in p.runs if 'literal' in r.text))
            with self.assertRaisesRegex(ValueError, 'Unsupported math command'):
                build.add_inline(doc.add_paragraph(), r'$\unsupported{x}$', decode_escapes=True)

    def test_v03_escaped_header_pipe_and_trailing_backslash(self):
        text = r'| \|Metric\| | Value |' + '\n|---|---|\n' + r'| path\\| $p_c$ |'
        headers, rows = list(build.md_tables(text))[0]
        self.assertEqual([build.cell_display(x) for x in headers], ['|Metric|', 'Value'])
        self.assertEqual(build.cell_display(rows[0][0]), 'path\\')

    def test_grayscale_copies_are_warned_and_never_embedded(self):
        with tempfile.TemporaryDirectory() as d:
            root = Path(d); self.fixture(root); out = root / 'output'
            extras = {}
            for family, prefix in [('main', 'fig'), ('supp', 's')]:
                folder = root / '02_figures/out/figures' / family
                # Reverse creation order also checks deterministic first-three examples.
                extras[folder] = [f'{prefix}{n:02}_grayscale.png' for n in range(1, 6)]
                for name in reversed(extras[folder]):
                    Image.new('L', (30, 30), 0).save(folder / name)
            manifest = build.build(root, output_dir=out, update_tracking=False)
            result = check.check_assembly(root, out)
            self.assertFalse(result['errors'], result)
            for report in (manifest, result):
                skipped = [w for w in report['warnings'] if w.startswith('Skipped non-numbered')]
                self.assertEqual(len(skipped), 2, skipped)
                for folder, names in extras.items():
                    entries = [w for w in skipped if str(folder) in w]
                    self.assertEqual(len(entries), 1, skipped)
                    self.assertIn('5 files', entries[0])
                    for name in names[:3]: self.assertIn(name, entries[0])
                    for name in names[3:]: self.assertNotIn(name, entries[0])
            for kind, source in [('Manuscript', 'main/fig1.png'),
                                 ('Supplementary', 'supp/s01.png')]:
                doc = Document(out / f'{root.name}_{kind}_v1.docx')
                self.assertEqual(len(doc.inline_shapes), 1)
                rid = doc.inline_shapes[0]._inline.graphic.graphicData.pic.blipFill.blip.embed
                self.assertEqual(doc.part.related_parts[rid].blob,
                                 (root / '02_figures/out/figures' / source).read_bytes())
            for source in ('main/fig1.png', 'supp/s01.png'):
                official = root / '02_figures/out/figures' / source
                original = official.read_bytes(); official.unlink()
                with self.assertRaisesRegex(ValueError, 'assets/captions mismatch'):
                    build.project_model(root)
                official.write_bytes(original)

    def test_arxiv_only_reference_and_integrity(self):
        with tempfile.TemporaryDirectory() as d:
            root = Path(d); self.fixture(root); out = root / 'output'
            intro = root / '04_manuscript/Introduction/Introduction_EN.md'
            intro.write_text('## Introduction\n\nA reported finding (arXiv: 2401.12345).', encoding='utf-8')
            meta = intro.with_name('Introduction_核查说明.md')
            meta.write_text('| 标识 | 第一作者 年份 | 期刊 | 题名 |\n|---|---|---|---|\n'
                            '| arXiv: 2401.12345 | A 2024 | arXiv | Preprint |', encoding='utf-8')
            manifest = build.build(root, output_dir=out, update_tracking=False)
            self.assertEqual([r['id'] for r in manifest['citation_candidates']], ['arXiv: 2401.12345'])
            self.assertFalse(check.check_assembly(root, out)['errors'])
            self.assertEqual(check.word_count('A finding (arXiv: 2401.12345)'), 2)
            changes = check.compare_text('(arXiv: 2401.12345)', '(DOI: 2401.12345)')
            self.assertTrue(any('Citation identifiers' in e for e in changes['errors']), changes)
            path = out / f'{root.name}_Manuscript_v1.docx'
            original = path.read_bytes()
            for mutation in ('added_reference','extra_body_identifier','lost_identifier'):
                with self.subTest(mutation=mutation):
                    path.write_bytes(original); doc=Document(path)
                    if mutation=='added_reference': doc.add_paragraph('References')
                    elif mutation=='extra_body_identifier': doc.add_paragraph('(arXiv: 2402.12345)')
                    else:
                        p=next(p for p in doc.paragraphs if 'arXiv:' in p.text)
                        p.text=p.text.replace('2401.12345','2401.99999')
                    doc.save(path)
                    self.assertTrue(check.check_assembly(root,out)['errors'])
            process=next((out/'_process').iterdir())
            self.assertIn('arXiv: 2401.12345',build.read(process/'参考文献候选清单.md'))
            self.assertNotIn('reference_count',check.check_assembly(root,out))

    def fixture(self, root):
        for section in (*build.SECTIONS,'Abstract'):
            folder=root/'04_manuscript'/section; folder.mkdir(parents=True)
            text=f'## {section}\n\nThe reported finding remains descriptive.'
            if section=='Introduction': text+=' (PMID: 22222222) then (PMID: 11111111).'
            if section=='Methods': text+=' The cohort (Table 1) and model (Fig. 1) are described.'
            if section=='Results': text+=' The result is reported (Table 2).'
            if section=='Abstract': text+='\n\n## Keywords\n\nImaging\n\n## Highlights\n\n- A reported finding.'
            (folder/f'{section}_EN.md').write_text(text,encoding='utf-8')
        leg=root/'04_manuscript/Legends';leg.mkdir()
        (leg/'Figure_Legends_EN.md').write_text('**Figure 1. Main.** Main caption.',encoding='utf-8')
        (leg/'Supplementary_Figure_Legends_EN.md').write_text('**Supplementary Figure S1. Extra.** Extra caption.',encoding='utf-8')
        (leg/'Table_Titles_Notes_EN.md').write_text('\n'.join(f'**{label}. Data.** Note.' for label in ('Table 1','Table 2','Supplementary Table S1')),encoding='utf-8')
        figs=root/'02_figures/out/figures'; (figs/'main').mkdir(parents=True);(figs/'supp').mkdir()
        Image.new('RGB',(200,100),'white').save(figs/'main/fig1.png')
        Image.new('RGB',(100,200),'gray').save(figs/'supp/s01.png')
        tab=root/'02_figures/out/tables';tab.mkdir()
        for label in ('T1','T2','ST01'):
            (tab/f'{label}.md').write_text('| Item | Value |\n|---|---|\n| A | 0.962 |',encoding='utf-8')
        meta=root/'04_manuscript/Introduction/Introduction_核查说明.md'
        meta.write_text('| 标识 | 第一作者 年份 | 期刊 | 题名 |\n|---|---|---|---|\n| PMID: 11111111 | B 2020 | Journal | Second |\n| PMID: 22222222 | A 2021 | Journal | First |',encoding='utf-8')

    def test_docx_content_and_deterministic_rerun(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d); self.fixture(p); out=p/'output'
            before={str(x.relative_to(p)):hashlib.sha256(x.read_bytes()).hexdigest() for x in p.rglob('*') if x.is_file()}
            first=build.build(p,output_dir=out,update_tracking=False)
            second=build.build(p,output_dir=out,update_tracking=False)
            self.assertEqual(first['outputs'],second['outputs'])
            result=check.check_assembly(p,out)
            self.assertFalse(result['errors'],result)
            for key,digest in before.items(): self.assertEqual(hashlib.sha256((p/key).read_bytes()).hexdigest(),digest)
            doc=Document(out/f'{p.name}_Manuscript_v1.docx')
            self.assertEqual(len(doc.inline_shapes),1);self.assertEqual(len(doc.tables),2)
            refs=[x.text for x in doc.paragraphs]
            self.assertNotIn('References',refs)
            candidates=build.read(next((out/'_process').iterdir())/'参考文献候选清单.md')
            self.assertLess(candidates.index('PMID: 22222222'),candidates.index('PMID: 11111111'))
            doc.tables[0].cell(1,1).text='0.963';doc.save(out/f'{p.name}_Manuscript_v1.docx')
            self.assertTrue(check.check_assembly(p,out)['errors'])

    def test_missing_caption_blocks_build(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d);self.fixture(p)
            (p/'04_manuscript/Legends/Figure_Legends_EN.md').write_text('**Figure 2. Wrong.** Caption.',encoding='utf-8')
            with self.assertRaises(ValueError): build.build(p,output_dir=p/'output')

    def test_omml_math_is_editable_and_unknown_command_fails(self):
        doc=build.new_document();par=doc.add_paragraph();build.add_inline(par,r'Rate \(3\times10^{-4}\) and \(W_{kc}\Delta\bar p_c\).')
        self.assertIn('m:sSup',par._p.xml);self.assertIn('m:sSub',par._p.xml);self.assertNotIn('\\times',par._p.xml)
        with self.assertRaises(ValueError):build.math_element(r'\unsupported{x}')

    def test_math_subscript_scope(self):
        from docx.oxml.ns import qn
        expr=build.math_element(r'p_c(W_{t,c}-W_{r,c})')
        sub=expr.find(qn('m:sSub')).find(qn('m:sub'))
        self.assertEqual(''.join(x.text or '' for x in sub.iter(qn('m:t'))),'c')
        bar=build.math_element(r'\bar p_c')
        accent=bar.find(qn('m:sSub')).find(qn('m:e')).find(qn('m:acc'))
        self.assertEqual(''.join(x.text or '' for x in accent.iter(qn('m:t'))),'p')

    def test_subsection_blank_lines(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d); self.fixture(p)
            (p/'04_manuscript/Methods/Methods_EN.md').write_text(NL.join([
                '## Methods', '', '## 2.1 Parent', '', '## 2.1.1 Child one', '', 'First child body (Table 1).', '',
                '## 2.1.2 Child two', '', 'Second child body (Fig. 1).', '', '## 2.2 Next', '', 'Next body (Table 2).']),
                encoding='utf-8')
            out = p/'output'; build.build(p, output_dir=out, update_tracking=False)
            doc = Document(out/f'{p.name}_Manuscript_v1.docx')
            text = [x.text.strip() for x in doc.paragraphs]
            def before(a, b):                      # b 前面紧邻的一段文本（资产块以空行包裹，中间可隔资产）
                return text[text.index(b) - 1]
            self.assertEqual(before('a', '2.1 Parent'), 'Methods')            # level1 -> child: no blank
            self.assertEqual(before('a', '2.1.1 Child one'), '2.1 Parent')    # parent -> child: no blank
            self.assertEqual(before('a', '2.1.2 Child two'), '')              # sibling subsection: blank right before
            self.assertEqual(before('a', '2.2 Next'), '')                     # subsection -> next section: blank right before
            # 资产在其所属子节末尾：Table 1 在 2.1.1 正文后、2.1.2 前；Fig. 1 在 2.1.2 正文后、2.2 前
            i_tab = text.index('First child body (Table 1).'); i_h2 = text.index('2.1.2 Child two')
            self.assertIn('Table 1.', ' '.join(text[i_tab:i_h2]))
            i_fig = text.index('Second child body (Fig. 1).'); i_next = text.index('2.2 Next')
            self.assertIn('Figure 1.', ' '.join(text[i_fig:i_next]))

    def test_abstract_runin_labels(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d); self.fixture(p)
            (p/'04_manuscript/Abstract/Abstract_EN.md').write_text(NL.join([
                '## Abstract', '', '### Background', '', 'BG body.', '', '### Methods', '', 'M body.', '',
                '### Results', '', 'R body.', '', '### Conclusions', '', 'C body.', '',
                '## Keywords', '', 'Imaging', '', '## Highlights', '', '- A reported finding.']), encoding='utf-8')
            out = p/'output'; build.build(p, output_dir=out, update_tracking=False)
            doc = Document(out/f'{p.name}_Manuscript_v1.docx')
            texts = [x.text.strip() for x in doc.paragraphs]
            for label, body in [('Background','BG body.'), ('Methods','M body.'), ('Results','R body.'), ('Conclusions','C body.')]:
                i = texts.index(f'{label}: {body}')                       # run-in: label+body in ONE paragraph
                runs = doc.paragraphs[i].runs
                self.assertEqual(runs[0].text, label + ':'); self.assertTrue(runs[0].bold)
                self.assertTrue(all(not r.bold for r in runs[1:] if r.text.strip()))
                self.assertEqual(runs[0].font.size.pt, 12)
            i0 = texts.index('Background: BG body.'); i3 = texts.index('Conclusions: C body.')
            self.assertEqual([t for t in texts[i0:i3+1] if not t], [])     # no blanks inside the abstract
            self.assertEqual(texts[texts.index('Abstract') + 1], 'Background: BG body.')
            self.assertEqual(texts[texts.index('Conclusions: C body.') + 1], '')   # blank before Keywords

    def test_tiff_preferred_with_png_fallback(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d); self.fixture(p)
            main = p/'02_figures/out/figures/main'
            m1 = build.build(p, output_dir=p/'o1', update_tracking=False)
            self.assertTrue(any('TIFF export missing for figure(s) 1' in w for w in m1['warnings']), m1['warnings'])
            Image.new('RGB', (200, 100), 'white').save(main/'fig1.tif', compression='tiff_lzw')
            m2 = build.build(p, output_dir=p/'o2', update_tracking=False)
            self.assertFalse(any('TIFF export missing for figure(s) 1 in' in w for w in m2['warnings']), m2['warnings'])
            doc = Document(p/'o2'/f'{p.name}_Manuscript_v1.docx')
            exts = sorted(str(r.target_part.partname).rsplit('.', 1)[-1].lower() for r in doc.part.rels.values() if 'image' in r.reltype)
            self.assertIn('tif', exts, exts)

if __name__ == '__main__':
    unittest.main(verbosity=2)
