"""Offline structural and before/after integrity checks; exit 0 clean, 1 errors.

Style counts are diagnostic only. They never authorize deletion of scientific evidence.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import date
import difflib
import json
from pathlib import Path
import re
import shutil

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
import build_manuscript as build

NUMBER = re.compile(r'(?<![\d.])(?:[+−-]?\d{1,3}(?:,\d{3})+(?:\.\d+)?|[+−-]?(?:\d+\.\d*|\.\d+|\d+))(?:[eE][+−-]?\d+)?')
QUALIFIERS = [r'\bsingle[- ]run\b', r'\bsingle evaluated seed\b', r'\bsingle[- ]seed\b',
              r'\bone (?:run|seed)\b', r'\bseed[- ]0\b', r'\bsubset\b', r'\bexploratory\b',
              r'\bdescriptive(?:ly)?\b', r'\bretrospective\b', r'\bnot conclusive\b',
              r'\binconclusive\b', r'\bcross(?:ing|ed|es)? zero\b', r'\bnot estimable\b',
              r'\b(?:includ(?:e|ed|es|ing)|contain(?:s|ed|ing)?) zero\b',
              r'\bnot executed\b', r'\bnot independent\b', r'\bobserved[- ]class\b',
              r'\bsubject[- ]separated\b', r'\bcluster[- ]bootstrap\b', r'\bpaired\b',
              r'\bone[- ]sided\b', r'\btwo[- ]sided\b', r'\bscan[- ]level\b',
              r'\bgroup[- ]level\b', r'\bconfidence intervals?\b', r'\bCI\b',
              r'\b(?:not|no|never|without|unestablished|subnominal|lower|failed|unresolved)\b',
              r'\blimitations?\b']
UNIT_TOKEN = r'(?:µm/px|μm/px|mmHg|mmol/L|mol/L|mg/dL|g/L|kg|mg|µg|μg|ng|pg|g|mL|ml|µL|μL|dL|L|mm|cm|µm|μm|nm|px|pixels?|scans?|eyes|subjects?|patients?|cohorts?|clusters?|seeds?|runs?|epochs?|seconds?|ms|Hz|GB|MB|%)'
UNITS = re.compile(r'(?<!\w)'+UNIT_TOKEN+r'(?!\w)|%')
MEASUREMENT = re.compile('('+NUMBER.pattern+r')\s*('+UNIT_TOKEN+r')(?!\w)')
FIG_ID = re.compile(r'\b(?:Fig(?:ure)?s?\.?|Tables?)\s+S?\d+[a-z]?(?:[–-][a-z])?', re.I)
OPERATORS = re.compile(r'<=|>=|≤|≥|<|>|≠|(?<![A-Za-z])=(?!=)')
DEFENSIVE = re.compile(r'\b(?:unfortunately|regrettably|we apologize|we must admit|merely|disappointing(?:ly)?|severely limited)\b', re.I)
AI_STYLE = re.compile(r'\b(?:delve|tapestry|pivotal|showcase|groundbreaking|it is (?:important|worth) to note|it is worth noting|in conclusion)\b|—|\bnot\b.{0,60}\bbut\b', re.I)
TARGETS = {'Methods': (2500, 3000), 'Results': (3500, 4000), 'Introduction': (0, 1000),
           'Discussion': (800, 1500), 'Abstract': (0, 250)}


def word_count(text):
    text = re.sub(r'^#{1,6}[^\n]*', '', text, flags=re.M)
    # Exclude reference and figure/table labels, not their surrounding prose.
    text = re.sub(r'\([^()]*\b(?:PMID|DOI|URL|arXiv):[^()]*\)', '', text, flags=re.I)
    text = re.sub(r'\((?:Supplementary\s+)?(?:Fig\.?|Figure|Table)[^()]*\)', '', text)
    text = build.plain(text)
    # Match pipeline's panel-label exclusion, including Chinese punctuation.
    # Word boundaries keep embedded function notation such as x(a) as content.
    text = re.sub(r'(?<!\w)(?:\([a-h](?:[–-][a-h])?\)|（[a-h](?:[–-][a-h])?）)'
                  r'(?!\w)[,，;；:：.。]?', ' ', text)
    return len(text.split())


def qualifier_counts(text):
    result = Counter()
    for pattern in QUALIFIERS:
        for m in re.finditer(pattern, text, re.I): result[m[0].lower()] += 1
    return result


def compare_text(before, after, protected=()):
    errors, warnings = [], []
    def same(name, a, b):
        if a != b:
            errors.append(f'{name} changed: removed={dict(a-b)}, added={dict(b-a)}')
    same('Numeric multiset', Counter(NUMBER.findall(before)), Counter(NUMBER.findall(after)))
    same('Units/analysis counts', Counter(UNITS.findall(before)), Counter(UNITS.findall(after)))
    same('Number/unit bindings', Counter(MEASUREMENT.findall(before)), Counter(MEASUREMENT.findall(after)))
    same('Citation identifiers', Counter(build.citation_ids(before, False)), Counter(build.citation_ids(after, False)))
    same('Figure/table labels', Counter(FIG_ID.findall(before)), Counter(FIG_ID.findall(after)))
    same('Statistical operators', Counter(OPERATORS.findall(before)), Counter(OPERATORS.findall(after)))
    # Formula, code, link targets and exact source terminology are not style-editing targets.
    same('Math expressions', Counter(m[0] for m in build.MATH_RE.finditer(before)),
         Counter(m[0] for m in build.MATH_RE.finditer(after)))
    same('Inline code', Counter(re.findall(r'`[^`]+`', before)), Counter(re.findall(r'`[^`]+`', after)))
    a, b = qualifier_counts(before), qualifier_counts(after)
    same('Protected qualifiers', a, b)
    literals = {s for s in protected if s and s in before}
    for literal in sorted(literals):
        if before.count(literal) != after.count(literal):
            errors.append('Protected literal changed: '+literal)
    wb, wa = word_count(before), word_count(after)
    if wb and not 0.9*wb <= wa <= 1.1*wb: errors.append(f'Word-count ±10% violated: {wb} -> {wa}')
    if not wb and wa: errors.append('Empty source gained text')
    return {'errors': errors, 'warnings': warnings, 'words_before': wb, 'words_after': wa,
            'qualifiers_before': dict(a), 'qualifiers_after': dict(b),
            'defensive_before': len(DEFENSIVE.findall(before)), 'defensive_after': len(DEFENSIVE.findall(after)),
            'ai_style_before': len(AI_STYLE.findall(before)), 'ai_style_after': len(AI_STYLE.findall(after))}


def paragraph_changes(before, after):
    a, b = build.blocks(before), build.blocks(after)
    matcher = difflib.SequenceMatcher(a=a, b=b, autojunk=False)
    return [{'before_start': i+1, 'after_start': j+1, 'before': '\n\n'.join(a[i:ie]),
             'after': '\n\n'.join(b[j:je])} for op, i, ie, j, je in matcher.get_opcodes() if op != 'equal']


def english_files(path):
    return {p.relative_to(path).as_posix(): p for p in path.rglob('*_EN.md')}


def compare_paths(before, after, protected=(), changes=None):
    before, after = Path(before), Path(after)
    if before.is_file() and after.is_file(): pairs = {before.name: (before, after)}
    elif before.is_dir() and after.is_dir():
        a, b = english_files(before), english_files(after)
        if set(a) != set(b): return {'errors': ['English file inventory changed: '+str(set(a)^set(b))], 'files': {}}
        if not a: return {'errors': ['No *_EN.md files found'], 'files': {}}
        pairs = {k: (a[k], b[k]) for k in sorted(a)}
    else: return {'errors': ['before/after must both be files or both be directories'], 'files': {}}
    results, errors = {}, []
    expected_logs = []
    for key, (a, b) in pairs.items():
        old, new = build.read(a), build.read(b)
        literals = protected.get(key, []) if isinstance(protected, dict) else protected
        result = compare_text(old, new, literals)
        for literal in literals:
            if literal and literal not in old: result['errors'].append('Protected literal not found in baseline: '+literal)
        section = Path(key).stem.removesuffix('_EN')
        if section in TARGETS:
            old_body, new_body = old, new
            if section == 'Abstract':
                old_body = re.split(r'^## Keywords', old, flags=re.M)[0]
                new_body = re.split(r'^## Keywords', new, flags=re.M)[0]
            nb, na = word_count(old_body), word_count(new_body)
            low, high = TARGETS[section]
            if not low <= na <= high:
                message = f'{section} absolute word target {low}–{high}: {na}'
                result['errors' if low <= nb <= high else 'warnings'].append(message)
            if section == 'Abstract' and nb and not 0.9*nb <= na <= 1.1*nb:
                result['errors'].append('Abstract body ±10% word band violated')
        if section in ('Figure_Legends', 'Supplementary_Figure_Legends'):
            old_caps, new_caps = build.captions(old, 'figure'), build.captions(new, 'figure')
            result['errors'] += match_ids(list(old_caps), list(new_caps), 'Polished captions')
            for cap in old_caps.keys() & new_caps.keys():
                nb, na = word_count(old_caps[cap]), word_count(new_caps[cap])
                if nb and not 0.9*nb <= na <= 1.1*nb: result['errors'].append('Figure '+cap+' ±10% word band violated')
                if na > 150:
                    result['errors' if nb <= 150 else 'warnings'].append(f'Figure {cap}: {na} words (>150)')
        result['changes'] = paragraph_changes(old, new)
        # Per-change numerical/qualifier checks prevent a same-file swap from hiding globally.
        for change in result['changes']:
            local = compare_text(change['before'], change['after'], literals)
            result['errors'] += ['Changed paragraph: '+e for e in local['errors'] if not e.startswith('Word-count')]
            expected_logs.append({'file': key, **change})
        results[key] = result
        errors += [key+': '+e for e in result['errors']]
    if changes is not None:
        raw = json.loads(build.read(changes))
        entries = raw.get('changes', []) if isinstance(raw, dict) else raw
        def signature(x):
            return (x.get('file'), x.get('before_start'), x.get('after_start'), x.get('before'), x.get('after'))
        if Counter(signature(x) for x in entries) != Counter(signature(x) for x in expected_logs):
            errors.append('Change log does not cover every changed paragraph exactly once')
        if any(not x.get('reason', '').strip() for x in entries): errors.append('Change log has empty reasons')
    return {'errors': errors, 'files': results, 'expected_changes': expected_logs,
            'boundary': 'Mechanical checks do not prove semantic equivalence; review every finding and adverse result.'}


def match_ids(assets, captions, name):
    return [] if Counter(assets) == Counter(captions) else [name+' assets/captions do not match']


def doc_paragraph_text(p):
    # Include editable Office Math text, which python-docx Paragraph.text omits.
    return ''.join(e.text or '' for e in p._p.iter() if e.tag in (qn('w:t'), qn('m:t')))


def rendered_text(text, decode_escapes=False):
    p = Paragraph(OxmlElement('w:p'), None); build.add_inline(p, text, decode_escapes=decode_escapes)
    return doc_paragraph_text(p)


def check_document(path, model, kind):
    errors, warnings = [], []
    doc = Document(path)
    texts = [doc_paragraph_text(p) for p in doc.paragraphs]
    joined = '\n'.join(texts)
    expected_figs = model['figures' if kind == 'main' else 'supp_figures']
    names = [shape._inline.docPr.get('name', '').removeprefix('Figure ') for shape in doc.inline_shapes]
    errors += match_ids(names, list(expected_figs), kind+' embedded figures')
    cap_key = 'figure_captions' if kind == 'main' else 'supp_captions'
    information = []
    body = list(doc.element.body)
    def check_caption(p, source, notes=False):
        end = 0 if notes else build.caption_title_end(source)
        expected = rendered_text(source[:end])
        expected += rendered_text(source[end:])
        if doc_paragraph_text(p) != expected:
            errors.append(kind+' approved caption/notes mismatch: '+source[:40])
        boundary = len(rendered_text(source[:end]))
        offset = 0
        for run in p._p.iter():
            if run.tag not in (qn('w:r'), qn('m:r')): continue
            text = ''.join(x.text or '' for x in run.iter() if x.tag in (qn('w:t'), qn('m:t')))
            flag = run.find('./'+qn('w:rPr')+'/'+qn('w:b'))
            bold = (flag.get(qn('w:val'), '1') not in ('0', 'false', 'off')
                    if flag is not None else p.style.font.bold)
            for pos in range(offset, offset+len(text)):
                if bool(bold) != (pos < boundary):
                    errors.append(kind+' caption bold extends beyond/misses title run: '+source[:40])
                    return
            offset += len(text)
    for key, caption in model[cap_key].items():
        target = rendered_text(caption)
        matches = [p for p in doc.paragraphs if doc_paragraph_text(p) == target]
        if len(matches) != 1:
            errors.append(f'{kind} figure {key} caption missing/duplicated'); continue
        p = matches[0]; check_caption(p, caption)
        shape = next((x for x in doc.inline_shapes if x._inline.docPr.get('name') == 'Figure '+key), None)
        if shape is not None:
            picture = shape._inline.getparent().getparent().getparent()
            if body.index(p._p) != body.index(picture)+1:
                errors.append(f'{kind} figure {key} caption must follow picture')
            if abs(shape.width.cm-14.6) > .1:
                errors.append(f'{kind} figure {key} width differs from 14.6 cm')
            from PIL import Image
            with Image.open(expected_figs[key]) as image:
                ratio = image.height/image.width
            if abs(shape.height/shape.width-ratio) > .002:
                errors.append(f'{kind} figure {key} aspect ratio changed')
        for paragraph in ([p, Paragraph(picture, doc)] if shape is not None else [p]):
            if paragraph.paragraph_format.keep_with_next or paragraph.paragraph_format.page_break_before:
                errors.append(f'{kind} figure {key} forced caption/page binding is obsolete')
    table_keys = sorted((k for k in model['tables'] if k.startswith('S') == (kind == 'supp')),
                        key=lambda k: int(k.removeprefix('S')))
    grouped = {k: [] for k in table_keys}; table_sequence = []; table_names = {}; layouts = []
    for table in doc.tables:
        previous = table._tbl.getprevious()
        if previous is None or previous.tag != qn('w:p'):
            errors.append(kind+' table title must be directly above table'); continue
        paragraph = Paragraph(previous, doc); caption = doc_paragraph_text(paragraph)
        match = re.match(r'(?:Supplementary\s+)?Table\s+(S?\d+)\.', caption)
        if not match or match[1] not in grouped:
            errors.append(kind+' unknown table caption: '+caption[:40]); continue
        key = match[1]; grouped[key].append(table); table_sequence.append(key); table_names[table._tbl] = key
        source = model['table_captions'][key]; end = build.caption_title_end(source)
        check_caption(paragraph, source[:end])
        if source[end:]:
            node = table._tbl.getnext()
            if node is None or node.tag != qn('w:p'):
                errors.append(f'Table {key}: notes missing after table')
            else:
                check_caption(Paragraph(node, doc), source[end:], notes=True)
        headers, rows = list(build.md_tables(build.read(model['tables'][key])))[0]
        columns = [[rendered_text(t, decode_escapes=True) for t in col] for col in zip(headers,*rows)]
        layout, widths = build.table_layout(columns, build.Cm(14.6).pt)
        layout = dict(table=key, **layout); layouts.append(layout)
        information.append(f'Table {key}: {layout["font_size_pt"]} pt; single table')
        if layout['word_wrap_columns']: warnings.append(build.table_wrap_warning(layout))
        actual = [[doc_paragraph_text(c.paragraphs[0]) for c in row.cells] for row in table.rows]
        expected = [[rendered_text(c, decode_escapes=True) for c in row] for row in [headers,*rows]]
        if actual != expected: errors.append(f'{kind} table {key} cell contents differ from source')
        sizes = [r.font.size.pt for row in table.rows for c in row.cells
                 for p in c.paragraphs for r in p.runs if r.text and r.font.size is not None]
        sizes += [int(x.get(qn('w:val')))/2 for x in table._tbl.xpath('.//m:r/w:rPr/w:sz')]
        if not sizes or set(sizes) != {layout['font_size_pt']}:
            errors.append(f'Table {key}: body font must use measured 8/7.5/7/6.5 pt step')
        if [c.width.twips for c in table.columns] != [w.twips for w in widths]:
            errors.append(f'Table {key}: column width differs from measured layout')
        if abs(sum(c.width.cm for c in table.columns)-14.6) > .02:
            errors.append(f'Table {key}: width differs from 14.6 cm')
        total = table._tbl.tblPr.find(qn('w:tblW'))
        if total is None or total.get(qn('w:type')) != 'dxa' or int(total.get(qn('w:w'))) != sum(w.twips for w in widths):
            errors.append(f'Table {key}: total width mismatch')
        for row in table.rows:
            for c,col in zip(row.cells, table.columns):
                if c.width.twips != col.width.twips: errors.append(f'Table {key}: grid/cell width mismatch')
                for paragraph in c.paragraphs:
                    if paragraph.paragraph_format.keep_with_next or paragraph.paragraph_format.keep_together:
                        errors.append(f'Table {key}: row must allow natural page flow')
                    wrap = paragraph._p.xpath('./w:pPr/w:wordWrap')
                    if not wrap or wrap[0].get(qn('w:val')) != '1':
                        errors.append(f'Table {key}: word-internal wrapping must be allowed')
        if table._tbl.xpath('.//w:cantSplit'): errors.append(f'Table {key}: rows must allow cross-page splitting')
        if not table.rows[0]._tr.xpath('./w:trPr/w:tblHeader'):
            errors.append(f'Table {key}: repeated header missing')
    for key,tables in grouped.items():
        if len(tables) != 1: errors.append(f'{kind} table {key}: expected single w:tbl, found {len(tables)}')
    if '(continued)' in joined: errors.append(kind+' continued table markers forbidden')
    if len(doc.sections) != 1 or doc.element.xpath('.//w:pgSz[@w:orient="landscape"]'):
        errors.append(kind+' must have one portrait section')
    for sec in doc.sections:
        if abs(sec.page_width.cm-21) > .01 or abs(sec.page_height.cm-29.7) > .01:
            errors.append(kind+' page must be A4 portrait')
        if any(abs(v.cm-3.2) > .01 for v in (sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin)):
            errors.append(kind+' page margins differ from 3.2 cm')
    if doc.element.xpath('.//w:lnNumType'): errors.append(kind+' line numbering must be absent')
    for rel in doc.part.rels.values():
        if rel.reltype.endswith(('/footer', '/header')):
            xml = rel.target_part.element.xml
            if 'PAGE' in xml: errors.append(kind+' PAGE field must be absent')
            if rel.target_part.element.xpath('.//w:t|.//w:fldSimple|.//w:instrText|.//w:drawing'):
                errors.append(kind+' header/footer content must be absent')
    for name, size in [('Normal', 12), ('Caption', 12)]:
        style = doc.styles[name]
        if style.font.name != 'Arial' or style.font.size.pt != size:
            errors.append(kind+' font mismatch: '+name)
        fonts = style.element.get_or_add_rPr().rFonts
        if fonts is not None and any('theme' in a.lower() for a in fonts.attrib):
            errors.append(kind+' Arial must not be overridden by theme: '+name)
        if style.paragraph_format.line_spacing != 1: errors.append(kind+' single spacing missing: '+name)
    for p in doc.paragraphs:
        for run in p.runs:
            if run.text and (run.font.name or p.style.font.name) != 'Arial':
                errors.append(kind+' English run font differs from Arial')
    if doc.element.xpath('.//w:br[@w:type="page"]'):
        errors.append(kind+' explicit page breaks must be absent')
    if kind == 'supp':
        if names != sorted(expected_figs, key=lambda k: int(k[1:])):
            errors.append('Supplementary figures out of numerical order')
        if table_sequence != sorted(table_sequence, key=lambda k: int(k[1:])):
            errors.append('Supplementary tables out of numerical order')
        if doc.tables and doc.inline_shapes:
            last_picture = doc.inline_shapes[-1]._inline.getparent().getparent().getparent()
            if body.index(last_picture) > body.index(doc.tables[0]._tbl):
                errors.append('Supplementary figures must precede tables')
    # Check body topology independently of paragraph styles. This also catches extra
    # prose, duplicate paragraphs, missing blanks and assets moved within a section.
    actual = []; paragraph_nodes = []
    for node in body:
        if node.tag == qn('w:sectPr'): continue
        if node.tag == qn('w:tbl'):
            actual.append(('table',table_names.get(node,'?'))); paragraph_nodes.append(None); continue
        p = Paragraph(node,doc)
        drawings = node.xpath('.//wp:docPr')
        if drawings: actual.append(('figure', drawings[0].get('name','').removeprefix('Figure ')))
        else: actual.append(('text',doc_paragraph_text(p)))
        paragraph_nodes.append(p)
    expected = []; heading_formats = {}
    def text(value, size=None):
        if size is not None: heading_formats[len(expected)] = size
        expected.append(('text',rendered_text(value)))
    def asset(family,key):
        if family == 'figure':
            expected.append(('figure',key)); text(model[cap_key][key])
        else:
            source = model['table_captions'][key]; end = build.caption_title_end(source)
            text(source[:end]); expected.append(('table',key))
            if source[end:]: text(source[end:])
    if kind == 'main':
        text('[Manuscript title]',16); text('[Authors]',10)
        text('[Affiliations]'); text('[Corresponding author]')
        for event in build.assembly_events(model):
            if event['type'] == 'blank': text('')
            elif event['type'] == 'heading': text(event['text'],event['size'])
            elif event['type'] == 'paragraph':
                text((event['runin'] + ': ' if event.get('runin') else '') + event['text'])
            else: asset(event['family'],event['key'])
    else:
        text('Supplementary material',16); text('')
        ordered = [('figure',k) for k in sorted(expected_figs,key=lambda k:int(k[1:]))]
        ordered += [('table',k) for k in table_keys]
        for family,key in ordered:
            asset(family,key); text('')
    if actual != expected:
        mismatch = next((i for i,(a,b) in enumerate(zip(actual,expected)) if a!=b), min(len(actual),len(expected)))
        errors.append(f'{kind} section-end first citation placement/source/blank order mismatch at node {mismatch}: '
                      f'actual={actual[mismatch:mismatch+1]}, expected={expected[mismatch:mismatch+1]}')
    for i,size in heading_formats.items():
        if i >= len(actual) or actual[i] != expected[i]: continue
        p = paragraph_nodes[i]
        if p.style.name != 'Normal' or not p.runs or any(
                r.font.size is None or r.font.size.pt != size or r.bold is not True or r.italic
                for r in p.runs if r.text):
            errors.append(kind+' heading must be Normal with direct bold/size: '+p.text)
    if doc.element.xpath('.//w:pStyle[starts-with(@w:val,"Heading") or starts-with(@w:val,"Title")]'):
        errors.append(kind+' Heading/Title paragraph styles forbidden')
    if doc.element.xpath('.//w:outlineLvl'): errors.append(kind+' outline levels forbidden')
    for p in doc.paragraphs:
        if p.style.element.xpath('./w:pPr/w:outlineLvl'): errors.append(kind+' inherited outline level forbidden')
        if not doc_paragraph_text(p) and not p._p.xpath('.//w:drawing'):
            if p.style.name != 'Normal' or any((r.font.name or p.style.font.name) != 'Arial' or
                    (r.font.size or p.style.font.size).pt != 12 or r.bold for r in p.runs):
                errors.append(kind+' blank paragraphs must be Normal Arial 12 pt regular')
    if any(re.fullmatch(r'\s*References\s*',t,re.I) for t in texts):
        errors.append(kind+' References must be absent')
    # No reference-list contract: identifiers remain part of the unchanged source text.
    return {'errors':errors,'warnings':warnings,'information':information,'table_layouts':layouts,
            'figures':len(doc.inline_shapes),'tables':len(doc.tables),'paragraphs':len(texts)}


def check_assembly(root, output=None, version=1, name=None):
    root = Path(root).resolve(); output = Path(output) if output else root/'05_final'; name = name or root.name
    model = build.project_model(root)
    results, errors, warnings = {}, [], list(model['warnings']); information = []
    for kind,suffix in [('main','Manuscript'),('supp','Supplementary')]:
        path = output/f'{name}_{suffix}_v{version}.docx'
        if not path.is_file(): errors.append('Missing output: '+str(path)); continue
        r = check_document(path,model,kind); results[kind] = r
        errors += r['errors']; warnings += r['warnings']; information += r['information']
    words = {}
    for section,text in model['texts'].items():
        if section == 'Abstract': text = re.split(r'^## Keywords',text,flags=re.M)[0]
        n = word_count(text); words[section] = n; low,high = TARGETS[section]
        if not low <= n <= high: warnings.append(f'Source {section} word count {n} outside {low}–{high}; do not auto-trim')
    for family in ('figure_captions','supp_captions'):
        for key,caption in model[family].items():
            n = word_count(caption)
            if n > 150: warnings.append(f'Source Figure {key} has {n} words (>150); author decision required')
    return {'errors':errors,'warnings':warnings,'information':information,'documents':results,'words':words,
            'submission_ready':False,
            'boundary':'No References is generated. Structural checks do not establish visual layout, submission readiness, clinical validity or author acceptance.'}


def apply_polish(before, after, changes, decisions, root, day, protected):
    """Archive all changed English originals before any overwrite; rollback on I/O error."""
    root, before, after = Path(root).resolve(), Path(before).resolve(), Path(after).resolve()
    if before != root/'04_manuscript' or not before.is_dir() or not after.is_dir():
        raise ValueError('--apply requires --before PROJECT/04_manuscript and a staged English tree')
    if not Path(decisions).is_file() or not build.read(decisions).strip(): raise ValueError('Nonempty decisions list required')
    if after == before or after.is_relative_to(before): raise ValueError('Stage outside 04_manuscript')
    result = compare_paths(before, after, protected, changes)
    if result['errors']: return result
    archive = root/'_archive'/day/'04_manuscript'
    files = english_files(before); candidates = english_files(after)
    if any(not p.resolve().is_relative_to(before) for p in files.values()):
        raise ValueError('English source symlink escapes manuscript root')
    changed = [key for key in files if files[key].read_bytes() != candidates[key].read_bytes()]
    backups = {}
    for key in changed:
        target = archive/key
        if target.exists() and target.read_bytes() != files[key].read_bytes():
            target = target.with_name(target.stem+'_'+build.sha(files[key])[:12]+target.suffix)
        target.parent.mkdir(parents=True, exist_ok=True)
        if not target.exists(): shutil.copyfile(files[key], target)
        if target.read_bytes() != files[key].read_bytes(): raise ValueError('Archive verification failed: '+key)
        backups[key] = target
    try:
        for key in changed:
            tmp = files[key].with_suffix('.md.tmp'); shutil.copyfile(candidates[key], tmp); tmp.replace(files[key])
    except OSError:
        for key in changed: shutil.copyfile(backups[key], files[key])
        raise
    result['applied'] = changed; result['archive'] = str(archive)
    return result


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--project-root', type=Path); p.add_argument('--output-dir', type=Path)
    p.add_argument('--version', type=int, default=1); p.add_argument('--name')
    p.add_argument('--before', type=Path); p.add_argument('--after', type=Path)
    p.add_argument('--protected-json', type=Path, help='JSON list of exact terms/evidence strings or a relative-file mapping')
    p.add_argument('--changes', type=Path, help='JSON paragraph changes with exact before/after and reasons')
    p.add_argument('--decisions', type=Path); p.add_argument('--apply', action='store_true')
    p.add_argument('--date', default=date.today().isoformat()); p.add_argument('--report', type=Path)
    args = p.parse_args()
    try:
        date.fromisoformat(args.date)
        protected = json.loads(build.read(args.protected_json)) if args.protected_json else []
        if bool(args.before) != bool(args.after): raise ValueError('Both --before and --after are required')
        if args.apply:
            if not all((args.before, args.after, args.changes, args.decisions, args.project_root, args.protected_json)):
                raise ValueError('--apply requires root/before/after/changes/decisions/protected-json')
            result = apply_polish(args.before, args.after, args.changes, args.decisions, args.project_root, args.date, protected)
        elif args.before:
            result = compare_paths(args.before, args.after, protected, args.changes)
        elif args.project_root:
            result = check_assembly(args.project_root, args.output_dir, args.version, args.name)
        else: raise ValueError('Choose --project-root assembly check or --before/--after comparison')
    except (ValueError, OSError, KeyError) as e:
        result = {'errors': [str(e)]}
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 1 if result['errors'] else 0


if __name__ == '__main__':
    raise SystemExit(main())
