"""Offline Markdown/PNG to manuscript DOCX. No network, model calls or source edits.

Use python -X utf8. See --help and references/assembly-rules.md.
"""
from __future__ import annotations

import argparse
from datetime import date, datetime
from functools import lru_cache
import hashlib
import io
import json
import math
from pathlib import Path
import re
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Length, Pt
from docx.text.paragraph import Paragraph
from PIL import Image, ImageFont

SECTIONS = ('Introduction', 'Methods', 'Results', 'Discussion')
CITE_RE = re.compile(r'\b(PMID|DOI|URL|arXiv):\s*(10\.[^\s;)\]}>]+|https?://[^\s;)\]}>]+|\d{4}\.\d{4,5}(?:v\d+)?|\d+)', re.I)
MATH_RE = re.compile(r'\\\((.*?)\\\)|\\\[(.*?)\\\]|\$([^$\n]+)\$', re.S)
MD_ESCAPE_RE = re.compile(r'\\([!"#$%&\'()*+,\-./:;<=>?@\[\]\\^_`{|}~])')
TABLE_MIN_FONT_PT = 6.5


def read(path):
    return Path(path).read_text(encoding='utf-8-sig')


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def plain(text):
    text = re.sub(r'\[([^\]]+)\]\([^\n]+?\)', r'\1', text)
    return text.replace('**', '').replace('`', '').strip()


def citation_ids(text, unique=True):
    found = [f'{"arXiv" if m[0].lower() == "arxiv" else m[0].upper()}: {m[1].rstrip(".,")}'
             for m in CITE_RE.findall(text)]
    return list(dict.fromkeys(found)) if unique else found


def table_cells(line):
    """A pipe is a separator only after an even number of backslashes."""
    cells, buf, escaped = [], [], False
    for char in line.strip():
        if char == '|' and not escaped:
            cells.append(''.join(buf).strip()); buf = []
        else:
            buf.append(char)
        escaped = not escaped if char == '\\' else False
    cells.append(''.join(buf).strip())
    if cells and not cells[0]: cells.pop(0)
    if cells and not cells[-1]: cells.pop()
    return cells


def md_tables(text):
    lines = text.splitlines()
    for i, line in enumerate(lines[:-1]):
        if line.strip().startswith('|') and re.fullmatch(r'[\s|:\-]+', lines[i+1]):
            header = table_cells(line)
            rows = []
            for row in lines[i+2:]:
                if not row.strip().startswith('|'):
                    break
                cells = table_cells(row)
                if len(cells) != len(header):
                    raise ValueError('Malformed Markdown table row: ' + row)
                rows.append(cells)
            yield header, rows


def blocks(text):
    """Markdown paragraphs, including a single-line list item as its own block."""
    out, buf = [], []
    for line in text.splitlines():
        if not line.strip() or re.match(r'^#{1,6}\s|^[-*]\s', line):
            if buf:
                out.append(' '.join(buf)); buf = []
            if line.strip():
                out.append(line.strip())
        else:
            buf.append(line.strip())
    if buf:
        out.append(' '.join(buf))
    return out


def captions(text, kind):
    label = r'(?:Supplementary\s+)?Figure' if kind == 'figure' else r'(?:Supplementary\s+)?Table'
    pattern = re.compile(r'^\s*(?:#{1,6}\s*)?(?:\*\*)?' + label + r'\s+(S?\d+)\.', re.M | re.I)
    hits = list(pattern.finditer(text))
    result = {}
    for i, hit in enumerate(hits):
        key = hit[1].upper()
        if key in result:
            raise ValueError(f'Duplicate {kind} caption {key}')
        result[key] = plain(text[hit.start():hits[i+1].start() if i+1 < len(hits) else len(text)])
    if not result:
        raise ValueError(f'No {kind} captions found')
    return result


def references(root, ids):
    """Only local verification tables. Never fabricate missing bibliographic fields."""
    catalog = {}
    for path in sorted((root/'04_manuscript').glob('*/*核查说明.md')):
        for headers, rows in md_tables(read(path)):
            headers = [plain(h) for h in headers]
            if not any(x in headers for x in ('标识', '正文标识')):
                continue
            id_col = next(i for i, x in enumerate(headers) if x in ('标识', '正文标识'))
            for row in rows:
                identifier = row[id_col]
                if 'URL:' in identifier:
                    link = re.search(r'\]\((https?://[^)]+)\)', identifier)
                    if link: identifier = 'URL: '+link[1]
                for key in citation_ids(identifier):
                    fields_by_name = {h:plain(c) for h,c in zip(headers,row)}
                    missing = [label for label,names in (
                        ('作者年份',('第一作者 年份',)), ('题名',('题名',)), ('期刊/出处',('期刊','期刊/出处')))
                        if not any(fields_by_name.get(n,'').strip('—- ') for n in names)]
                    full = not missing
                    if full:
                        def val(*names):
                            return next((plain(row[headers.index(n)]) for n in names if n in headers), '')
                        fields = [val('第一作者 年份'), val('题名').rstrip('.'), val('期刊', '期刊/出处')]
                        display = '. '.join(x for x in fields if x) + '. ' + key
                    else:
                        display = plain(row[0]) + '. ' + key
                    record = {'id': key, 'text': display, 'complete': full,
                              'source': str(path.relative_to(root)), 'row': row,
                              'fields': fields_by_name, 'missing_fields': missing,
                              'unverified': any('暂未核验' in c for c in row)}
                    if key not in catalog or (full and not catalog[key]['complete']):
                        catalog[key] = record
    return [catalog.get(key, {'id': key, 'text': key, 'complete': False,
                             'source': None, 'unverified': True,
                             'fields': {}, 'missing_fields': ['作者年份','题名','期刊/出处']}) for key in ids]


def figure_assets(folder, prefix, warnings):
    """组图资产：优先 TIFF（投稿与期刊要求），PNG 为回退并出 WARN；只认精确编号文件名。"""
    assets = {}
    skipped = []
    numbered = set()
    for path in sorted(folder.glob('*')):
        if not path.is_file():
            continue
        stem, suffix = path.stem, path.suffix.lower()
        if suffix not in ('.png', '.tif', '.tiff'):
            continue
        if not re.fullmatch(prefix + r'[0-9]+', stem):
            skipped.append(path.name)
            continue
        try:
            number = int(stem[len(prefix):])
        except ValueError:
            warnings.append('Skipped unparseable figure number: ' + path.name)
            continue
        key = ('S' if prefix == 's' else '') + str(number)
        numbered.add(key)
        if suffix in ('.tif', '.tiff'):
            assets[key] = path                      # TIFF wins regardless of sort order
        elif key not in assets:
            assets[key] = path
    fallback = sorted(k for k, p in assets.items() if p.suffix.lower() == '.png')
    if fallback:
        warnings.append('TIFF export missing for figure(s) ' + ', '.join(fallback)
                        + f' in {folder}: PNG embedded (regenerate the figure with TIFF export)')
    if skipped:
        warnings.append(f'Skipped non-numbered figure candidates in {folder}: '
                        f'{len(skipped)} files; examples: ' + ', '.join(skipped[:3]))
    return assets


def project_model(root):
    root = Path(root).resolve()
    manuscript = root/'04_manuscript'
    texts = {s: read(manuscript/s/f'{s}_EN.md') for s in (*SECTIONS, 'Abstract')}
    leg = manuscript/'Legends'
    figs = captions(read(leg/'Figure_Legends_EN.md'), 'figure')
    supp = captions(read(leg/'Supplementary_Figure_Legends_EN.md'), 'figure')
    tabs = captions(read(leg/'Table_Titles_Notes_EN.md'), 'table')
    # Discover counts from source assets; require a consecutive numbering sequence.
    warnings = []
    main_assets = figure_assets(root/'02_figures/out/figures/main', 'fig', warnings)
    first = root/'03_outline/2_Methods/2.2_Architecture/figures/Fig1.tif'
    if first.is_file():
        main_assets['1'] = first
    supp_assets = figure_assets(root/'02_figures/out/figures/supp', 's', warnings)
    table_assets = {}
    for p in (root/'02_figures/out/tables').glob('*.md'):
        m = re.fullmatch(r'(S?T)(\d+)', p.stem, re.I)
        if m:
            key = ('S' if m[1].upper() == 'ST' else '')+str(int(m[2]))
            table_assets[key] = p
    for name, assets, caps in [('main figures', main_assets, figs), ('supp figures', supp_assets, supp),
                               ('tables', table_assets, tabs)]:
        if set(assets) != set(caps):
            raise ValueError(f'{name}: assets/captions mismatch: {set(assets)^set(caps)}')
        for prefix in ('', 'S'):
            nums = sorted(int(x.removeprefix('S')) for x in assets if x.startswith('S') == bool(prefix))
            if nums and nums != list(range(1, max(nums)+1)):
                raise ValueError(f'{name}: nonconsecutive numbering')
    ids = citation_ids('\n'.join(texts[s] for s in ('Abstract', *SECTIONS)) + '\n' +
                       '\n'.join(list(figs.values())+list(supp.values())+list(tabs.values())))
    return dict(root=root, texts=texts, figures=main_assets, supp_figures=supp_assets,
                figure_captions=figs, supp_captions=supp, table_captions=tabs,
                tables=table_assets, references=references(root, ids), warnings=warnings)


def math_element(expr):
    """Small, explicit TeX -> editable OMML subset. Unknown commands fail closed."""
    symbols = {'ell': 'ℓ', 'times': '×', 'beta': 'β', 'alpha': 'α', 'Delta': 'Δ',
               'in': '∈', 'mu': 'μ', 'sigma': 'σ', 'leq': '≤', 'geq': '≥', 'pm': '±'}
    tokens = re.findall(r'\\[A-Za-z]+|\\[{}]|.', expr, re.S)
    idx = 0
    def node(text):
        r = OxmlElement('m:r'); t = OxmlElement('m:t'); t.text = text; r.append(t)
        return r
    def atom():
        nonlocal idx
        while idx < len(tokens) and tokens[idx].isspace(): idx += 1
        if idx >= len(tokens):
            raise ValueError('Incomplete math expression')
        token = tokens[idx]; idx += 1
        if token == '{':
            return seq(True)
        if token == r'\bar':
            a = OxmlElement('m:acc'); pr = OxmlElement('m:accPr')
            char = OxmlElement('m:chr'); char.set(qn('m:val'), '̄'); pr.append(char); a.append(pr)
            e = OxmlElement('m:e'); e.extend(atom()); a.append(e); return [a]
        if token.startswith('\\'):
            cmd = token[1:]
            if cmd in ('{', '}'): return [node(cmd)]
            if cmd not in symbols: raise ValueError('Unsupported math command: '+token)
            token = symbols[cmd]
        if token in ('}', '_', '^'): raise ValueError('Malformed math expression')
        # Ordinary chunks split so the preceding character receives sub/superscripts.
        return [node(c) for c in token]
    def seq(group=False):
        nonlocal idx
        out = []
        while idx < len(tokens):
            if tokens[idx].isspace():
                idx += 1; continue
            if tokens[idx] == '}':
                if not group: raise ValueError('Unmatched math brace')
                idx += 1; return out
            if tokens[idx] in ('_', '^'):
                op = tokens[idx]; idx += 1
                if not out: raise ValueError('Math script without base')
                base = out.pop(); obj = OxmlElement('m:sSub' if op == '_' else 'm:sSup')
                e = OxmlElement('m:e'); e.append(base); obj.append(e)
                sub = OxmlElement('m:sub' if op == '_' else 'm:sup'); sub.extend(atom()); obj.append(sub)
                out.append(obj)
            else:
                out.extend(atom())
        if group: raise ValueError('Unclosed math brace')
        return out
    omath = OxmlElement('m:oMath'); omath.extend(seq()); return omath


def add_inline(paragraph, text, decode_escapes=False):
    def add_text(span):
        # Protect escaped punctuation before interpreting Markdown formatting.
        # Math has already been separated and never enters this decoder.
        literals = []
        if decode_escapes:
            def protect(match):
                literals.append(match[1]); return f'\ue000{len(literals)-1}\ue001'
            span = MD_ESCAPE_RE.sub(protect, span)
        span = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1', span)
        for part in re.split(r'(\*\*.*?\*\*|`[^`]+`)', span, flags=re.S):
            if not part: continue
            bold = part.startswith('**') and part.endswith('**')
            content = part[2:-2] if bold else part.strip('`')
            if decode_escapes:
                content = re.sub(r'\ue000(\d+)\ue001', lambda m: literals[int(m[1])], content)
            run = paragraph.add_run(content); run.bold = bold
    start = 0
    for match in MATH_RE.finditer(text):
        add_text(text[start:match.start()])
        equation = math_element(next(x for x in match.groups() if x is not None))
        for run in equation.iter(qn('m:r')):
            props = OxmlElement('w:rPr'); fonts = OxmlElement('w:rFonts')
            fonts.set(qn('w:ascii'), 'Arial'); fonts.set(qn('w:hAnsi'), 'Arial')
            props.append(fonts); run.insert(0, props)
        paragraph._p.append(equation)
        start = match.end()
    add_text(text[start:])


@lru_cache(maxsize=8)
def layout_font(size, bold=False):
    # Measure at 10x resolution in points. Fallback deliberately overestimates.
    for name in (('C:/Windows/Fonts/arialbd.ttf' if bold else 'C:/Windows/Fonts/arial.ttf'),
                 ('DejaVuSans-Bold.ttf' if bold else 'DejaVuSans.ttf')):
        try: return ImageFont.truetype(name, round(size*10))
        except OSError: pass
    return None


def text_width(text, size, bold=False):
    font = layout_font(size, bold)
    return font.getlength(text)/10 if font else len(text)*size


def cell_display(text):
    p = Paragraph(OxmlElement('w:p'), None)
    add_inline(p, text, decode_escapes=True)
    return ''.join(e.text or '' for e in p._p.iter() if e.tag in (qn('w:t'), qn('m:t')))


def new_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(3.2)
    for name, size, bold in (('Normal', 12, False), ('Title', 16, True),
                             ('Heading 1', 15, True), ('Heading 2', 12, True),
                             ('Heading 3', 12, True), ('Caption', 12, False)):
        style = doc.styles[name]
        style.font.name = 'Arial'; style.font.size = Pt(size); style.font.bold = bold
        style.font.italic = False
        fonts = style.element.get_or_add_rPr().rFonts
        for attr in list(fonts.attrib):
            if 'theme' in attr.lower(): del fonts.attrib[attr]
        for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'): fonts.set(qn('w:'+attr), 'Arial')
        style.font.color.rgb = __import__('docx').shared.RGBColor(0, 0, 0)
        style.paragraph_format.line_spacing = 1
        style.paragraph_format.space_before = style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.keep_with_next = name.startswith('Heading')
        style.paragraph_format.keep_together = False
    doc.styles['Normal'].paragraph_format.widow_control = True
    doc.core_properties.author = ''; doc.core_properties.last_modified_by = ''
    doc.core_properties.created = doc.core_properties.modified = datetime(2000, 1, 1)
    return doc


def add_heading(doc, text, size=12):
    """Plain Normal paragraph with direct formatting; no navigation outline."""
    p = doc.add_paragraph(style='Normal')
    add_inline(p, text)
    for run in p.runs:
        run.font.name = 'Arial'; run.font.size = Pt(size); run.bold = True
        run.italic = False
    return p


def add_blank(doc):
    p = doc.add_paragraph(style='Normal')
    run = p.add_run(''); run.font.name = 'Arial'; run.font.size = Pt(12); run.bold = False
    return p


def source_events(texts):
    """Semantic heading levels survive independently of Word paragraph styles."""
    events = []
    for section in ('Abstract', *SECTIONS):
        events.append(dict(type='heading', text=section, level=1, size=15))
        runin = None
        for block in blocks(texts[section]):
            m = re.match(r'^(#{1,6})\s+(.*)', block)
            if m:
                if re.fullmatch(r'(?:\d+\.?\s*)?'+section, m[2], re.I): continue
                if section == 'Abstract' and m[2] not in ('Keywords', 'Highlights'):
                    runin = m[2].rstrip(':').strip()   # run-in label: bold "Background:" + same-paragraph body
                    continue
                numeric = re.match(r'^(\d+(?:\.\d+)+)\b', m[2])
                level = len(numeric[1].split('.')) if numeric else max(2, len(m[1]))
                if m[2] in ('Keywords', 'Highlights'): level = 1
                events.append(dict(type='heading', text=m[2], level=level,
                                   size=15 if level == 1 else 12))
            else:
                event = dict(type='paragraph', text='• '+block[2:] if re.match(r'^[-*]\s', block) else block)
                if runin is not None: event['runin'] = runin; runin = None
                events.append(event)
    return events


def assembly_events(model):
    """First body mention selects its nearest heading, ending at next level <= L."""
    events = source_events(model['texts']); insertions = {}; seen = set(); heading = 0
    for i, event in enumerate(events):
        if event['type'] == 'heading': heading = i; continue
        for family, key in main_citations(event['text']):
            if (family, key) in seen: continue
            assets = model['tables' if family == 'table' else 'figures']
            if key not in assets: raise ValueError(f'Main {family} {key} cited without asset')
            seen.add((family, key))
            level = events[heading]['level']
            end = next((j for j in range(heading+1, len(events))
                        if events[j]['type'] == 'heading' and events[j]['level'] <= level), len(events))
            insertions.setdefault(end, []).append(dict(type='asset', family=family, key=key))
    required = {('table', x) for x in model['tables'] if not x.startswith('S')}
    required |= {('figure', x) for x in model['figures']}
    if seen != required: raise ValueError('Main assets not cited in body: '+str(required-seen))
    result = []
    def blank():
        if not result or result[-1]['type'] != 'blank': result.append(dict(type='blank'))
    for i in range(len(events)+1):
        for event in insertions.get(i, []):
            blank(); result.append(event); blank()
        if i == len(events): break
        event = events[i]
        if event['type'] == 'heading':
            prev = next((e for e in reversed(result) if e['type'] != 'blank'), None)
            parent_child = (prev is not None and prev['type'] == 'heading'
                            and prev['level'] == event['level'] - 1)
            # Blank line before every heading except a child heading directly under its parent heading.
            if event['level'] == 1 or not parent_child:
                blank()
        result.append(event)
    return result


def add_markdown(doc, text, skip_top=True, on_paragraph=None):
    for block in blocks(text):
        m = re.match(r'^(#{1,6})\s+(.*)', block)
        if m:
            # Source sections use # / ## as the top heading; only the section name is omitted.
            if skip_top and re.fullmatch(r'(?:\d+\.?\s*)?(?:Introduction|Methods|Results|Discussion|Abstract)', m[2], re.I):
                continue
            add_heading(doc, m[2], 15 if m[2] in ('Keywords', 'Highlights') else 12)
        else:
            p = doc.add_paragraph()
            if re.match(r'^[-*]\s', block): block = '• '+block[2:]
            add_inline(p, block)
            if on_paragraph: on_paragraph(block)


def table_plan(columns, available):
    """One table; reduce font to 6.5 pt, then permit word-internal wrapping."""
    budget = math.floor(available*20 + 1e-7)
    for size in (8, 7.5, 7, TABLE_MIN_FONT_PT):
        word, content = column_width_requirements(columns, size)
        if sum(word) <= budget:
            return dict(size=size)
    return dict(size=TABLE_MIN_FONT_PT)


def proportional_widths(columns, available, size):
    """Full 14.6 cm grid; keep measured word floors, distribute remainder by content."""
    word, content = column_width_requirements(columns, size)
    budget = round(available*20)
    if sum(word) > budget:
        # Keep short identifier/value columns intact before reducing wide columns.
        padding = 256
        if budget <= padding*len(columns): raise ValueError('Too many columns for cell padding')
        widths = [0]*len(columns); remaining = budget; pending = list(range(len(columns)))
        for i in sorted(pending, key=lambda i: word[i]):
            if word[i] > remaining/len(pending): break
            widths[i] = word[i]; remaining -= word[i]; pending.remove(i)
        weights = {i:max(1,word[i]-padding) for i in pending}
        total = sum(weights.values()); extra = remaining-padding*len(pending)
        shares = {i:extra*weights[i]/total for i in pending}
        for i in pending: widths[i] = padding+math.floor(shares[i])
        for i in sorted(pending, key=lambda i: (-(shares[i] % 1), i))[:budget-sum(widths)]:
            widths[i] += 1
        return [Pt(w/20) for w in widths]
    extra = budget-sum(word); total = sum(content)
    shares = [extra*c/total for c in content]
    widths = [w+math.floor(x) for w, x in zip(word, shares)]
    for i in sorted(range(len(widths)), key=lambda i: (-(shares[i] % 1), i))[:budget-sum(widths)]:
        widths[i] += 1
    return [Pt(w/20) for w in widths]


def word_tokens(text):
    """Whitespace boundaries except explicitly nonbreaking spaces; keep punctuation."""
    return re.findall(r'[^\s\u00a0\u202f]+(?:[\u00a0\u202f][^\s\u00a0\u202f]+)*', text)


def column_width_requirements(columns, size):
    """Word/content widths in twips, including headers, 8% glyph slack and margins."""
    word = [math.ceil((max((text_width(w, size) for t in col for w in word_tokens(t)),
                          default=0)*1.08 + 12.8)*20) for col in columns]
    content = [max(w, math.ceil((max((text_width(t, size) for t in col), default=0)*1.08
                                + 12.8)*20)) for col, w in zip(columns, word)]
    return word, content


def caption_title_end(caption):
    """Label plus first title sentence; panel labels always start regular text."""
    label = re.match(r'(?:Supplementary\s+)?(?:Figure|Fig\.?|Table)\s+S?\d+\.', caption, re.I)
    start = label.end() if label else 0
    tail = caption[start:]
    stop = re.search(r'[.!?。！？](?=\s|$)|(?<!\w)\([a-h](?:[–-][a-h])?\)', tail)
    if stop:
        return start + (stop.start() if stop[0].startswith('(') else stop.end())
    return len(caption)


def add_caption(doc, caption, title_only=False):
    p = doc.add_paragraph(style='Caption')
    end = caption_title_end(caption)
    for text, bold in ((caption[:end], True), ('' if title_only else caption[end:], False)):
        if not text: continue
        before = len(p.runs); children = len(p._p); add_inline(p, text)
        for run in p.runs[before:]: run.bold = bold
        for child in list(p._p)[children:]:
            for math_run in child.iter(qn('m:r')):
                props = math_run.find(qn('w:rPr'))
                flag = OxmlElement('w:b'); flag.set(qn('w:val'), '1' if bold else '0')
                props.append(flag)
    p.paragraph_format.keep_with_next = False
    p.paragraph_format.keep_together = False
    p.paragraph_format.page_break_before = False
    return p


def main_citations(text):
    """Body paragraphs only. Narrative and parenthetical mentions, excluding S numbers."""
    pattern = r'\b(?:Fig(?:ure)?s?\.?|Tables?)\s+(\d+)(?!\d)'
    for m in re.finditer(pattern, text, re.I):
        if re.search(r'Supplementary\s*$', text[:m.start()], re.I): continue
        kind = 'table' if m[0].lower().startswith('table') else 'figure'
        yield kind, m[1]


def table_layout(columns, available):
    plan = table_plan(columns, available)
    widths = proportional_widths(columns, available, plan['size'])
    wrap = []
    for i, (col, width) in enumerate(zip(columns, widths)):
        samples = list(dict.fromkeys(w for t in col for w in word_tokens(t)
            if (text_width(w, plan['size'])*1.08+12.8)*20 > width.twips))
        if samples: wrap.append(dict(column=i+1, header=col[0], samples=samples))
    return dict(font_size_pt=plan['size'], word_wrap_columns=wrap), widths


def table_wrap_warning(layout):
    details = '; '.join(f'column {c["column"]} ({c["header"]}): '+', '.join(c['samples'][:3])
                        for c in layout['word_wrap_columns'])
    return f'Table {layout["table"]}: estimated word-internal wrapping at 6.5 pt; '+details


def add_table(doc, key, model, warnings=None, information=None, layouts=None):
    parsed = list(md_tables(read(model['tables'][key])))
    if len(parsed) != 1: raise ValueError(f'Table {key}: expected one Markdown table')
    headers, rows = parsed[0]
    columns = [[cell_display(t) for t in col] for col in zip(headers, *rows)]
    layout, widths = table_layout(columns, Cm(14.6).pt)
    layout = dict(table=key, **layout)
    if layouts is not None: layouts.append(layout)
    if information is not None: information.append(f'Table {key}: {layout["font_size_pt"]} pt; single table')
    if layout['word_wrap_columns'] and warnings is not None:
        warnings.append(table_wrap_warning(layout))
    caption = model['table_captions'][key]; end = caption_title_end(caption)
    add_caption(doc, caption, title_only=True)
    table = doc.add_table(rows=1, cols=len(headers)); table.style = 'Table Grid'
    table.autofit = False
    tbl_width = table._tbl.tblPr.find(qn('w:tblW')); tbl_width.set(qn('w:type'), 'dxa')
    tbl_width.set(qn('w:w'), str(sum(w.twips for w in widths)))
    for column, width in zip(table.columns, widths): column.width = width
    repeat = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row_index, values in enumerate([headers, *rows]):
        cells = table.rows[0].cells if row_index == 0 else table.add_row().cells
        # No cantSplit and no paragraph keepNext: allow natural cross-page rows.
        for cell, value, width in zip(cells, values, widths):
            cell.width = width; paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.line_spacing = 1
            paragraph.paragraph_format.keep_with_next = False
            paragraph.paragraph_format.keep_together = False
            wrap = OxmlElement('w:wordWrap'); wrap.set(qn('w:val'), '1')
            paragraph._p.get_or_add_pPr().append(wrap)
            add_inline(paragraph, value, decode_escapes=True)
            for run in paragraph.runs:
                run.font.size = Pt(layout['font_size_pt']); run.font.name = 'Arial'
            for math_run in paragraph._p.iter(qn('m:r')):
                props = math_run.find(qn('w:rPr')); sz = OxmlElement('w:sz')
                sz.set(qn('w:val'), str(round(layout['font_size_pt']*2))); props.append(sz)
    if caption[end:]:
        p = doc.add_paragraph(style='Caption'); add_inline(p, caption[end:])
        for run in p.runs: run.bold = False
    return layout


def add_figure(doc, key, path, caption, process, warnings=None):
    path = Path(path)   # TIFF 直接嵌入（python-docx 支持 TIFF；不再转换 PNG）
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.keep_with_next = False
    p.paragraph_format.keep_together = False
    p.paragraph_format.page_break_before = False
    drawing = p.add_run().add_picture(str(path), width=Cm(14.6))
    drawing._inline.docPr.set('name', 'Figure '+key)
    drawing._inline.docPr.set('descr', caption.split('. ', 1)[0])
    add_caption(doc, caption)


def stable_save(doc, path):
    """Normalize ZIP ordering/timestamps so unchanged inputs yield identical bytes."""
    raw = io.BytesIO(); doc.save(raw)
    result = io.BytesIO()
    with zipfile.ZipFile(raw) as src, zipfile.ZipFile(result, 'w', zipfile.ZIP_DEFLATED) as dst:
        for name in sorted(src.namelist()):
            z = zipfile.ZipInfo(name, (1980, 1, 1, 0, 0, 0)); z.compress_type = zipfile.ZIP_DEFLATED
            data = src.read(name)
            dst.writestr(z, data)
    return result.getvalue()


def archive_write(path, data, root, day):
    path = Path(path)
    if path.exists():
        old = path.read_bytes()
        if old == data: return
        try: rel = path.relative_to(root)
        except ValueError: rel = Path('external-output')/path.name
        archived = root/'_archive'/day/rel.parent/(path.stem+'_'+hashlib.sha256(old).hexdigest()[:12]+path.suffix)
        archived.parent.mkdir(parents=True, exist_ok=True)
        if not archived.exists(): archived.write_bytes(old)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix+'.tmp'); tmp.write_bytes(data); tmp.replace(path)


def tracking(root, version, day):
    project = root/'PROJECT.md'
    if not project.exists(): return
    text = read(project)
    row = f'| 11 | 整稿组装与英文润色 | medical-sci-finalize | 已组装，待润色/作者验收 | 05_final/ | v{version} | {day} | 组装检查不代表科学核验或作者确认 |'
    if re.search(r'^\| 11 \|', text, re.M):
        text = re.sub(r'^\| 11 \|.*$', row, text, flags=re.M)
    else:
        matches = list(re.finditer(r'^\| 10 \|.*$', text, re.M))
        if not matches: raise ValueError('PROJECT.md lacks stage 10 row; cannot safely update tracking')
        m = matches[-1]; text = text[:m.end()]+'\n'+row+text[m.end():]
    archive_write(project, text.encode('utf-8'), root, day)
    decision = root/'04_manuscript/作者决定.md'
    content = read(decision) if decision.exists() else '# 作者决定\n'
    entry = '- Finalize 固定约定：两份 Word；主图主表置于首次引用所属节末，单节 A4 纵向，Arial 单倍行距，单表缩字至 6.5 pt；不输出 References；保守去防御，保留全部证据限定；英文归档后覆盖，中文不动。'
    if '## Finalize' not in content:
        content += '\n\n## Finalize\n\n'+entry+'\n'
        archive_write(decision, content.encode('utf-8'), root, day)


def write_process_lists(process, candidates, layouts):
    lines = ['# 参考文献候选清单', '', '按源正文首次出现顺序；随后收录仅在图表注中出现的标识。仅用本地字段，不联网核验。', '']
    for i, record in enumerate(candidates, 1):
        gaps = []
        if not record['complete']: gaps.append('书目字段不全：'+', '.join(record['missing_fields']))
        if record['source'] is None: gaps.append('核查表无记录')
        elif record['unverified']: gaps.append('明示暂未核验')
        lines += [f'{i}. {record["id"]}', f'   - 已有字段：{record["text"]}',
                  '   - 缺口：'+('；'.join(gaps) or '本地字段齐备；未作新增核验'),
                  '   - 来源：'+str(record['source'] or '无')]
        if record.get('row'): lines.append('   - 核查表原字段：'+' | '.join(record['row']))
        for field,value in record['fields'].items(): lines.append(f'   - {field}：{value or "空"}')
    (process/'参考文献候选清单.md').write_text('\n'.join(lines)+'\n', encoding='utf-8')
    lines = ['# 表格断行清单', '', '字号与断词为字体测量估算；实际分页和可读性需 Word 复核。', '']
    for layout in layouts:
        lines += [f'## Table {layout["table"]} — {layout["font_size_pt"]} pt', '']
        if not layout['word_wrap_columns']: lines.append('无估算词内断行。')
        for col in layout['word_wrap_columns']:
            lines.append(f'- 列 {col["column"]}（{col["header"]}）：'+', '.join(col['samples']))
        lines.append('')
    (process/'表格断行清单.md').write_text('\n'.join(lines)+'\n', encoding='utf-8')


def build(root, version=1, kind='both', output_dir=None, day=None, name=None, update_tracking=True):
    root = Path(root).resolve(); day = day or date.today().isoformat()
    date.fromisoformat(day)
    if version < 1: raise ValueError('version must be positive')
    pipeline = (root/'PROJECT.md').is_file()
    if pipeline and output_dir is not None and Path(output_dir).resolve() != root/'05_final':
        raise ValueError('Pipeline project output must be 05_final/')
    if not pipeline and output_dir is None: raise ValueError('No PROJECT.md: --output-dir is required')
    output = root/'05_final' if pipeline else Path(output_dir).resolve()
    if output == root or output.is_relative_to(root/'02_figures') or output.is_relative_to(root/'03_outline') or output.is_relative_to(root/'04_manuscript'):
        raise ValueError('Unsafe output directory')
    name = name or root.name
    if re.search(r'[<>:"/\\|?*]', name): raise ValueError('Invalid project name')
    model = project_model(root)
    process = root/'_process/Finalize'/f'{day}_assembly_v{version}' if pipeline else output/'_process'/f'{day}_assembly_v{version}'
    process.mkdir(parents=True, exist_ok=True)
    warnings = list(model['warnings']); information = []
    layouts = []; docs = {}
    if kind in ('both', 'main'):
        doc = new_document()
        add_heading(doc, '[Manuscript title]', 16)
        add_heading(doc, '[Authors]', 10)
        doc.add_paragraph('[Affiliations]'); doc.add_paragraph('[Corresponding author]')
        for event in assembly_events(model):
            if event['type'] == 'blank': add_blank(doc)
            elif event['type'] == 'heading': add_heading(doc, event['text'], event['size'])
            elif event['type'] == 'paragraph':
                p = doc.add_paragraph()
                if event.get('runin'):
                    label = p.add_run(event['runin'] + ':')
                    label.font.name = 'Arial'; label.font.size = Pt(12); label.bold = True; label.italic = False
                    gap = p.add_run(' ')
                    gap.font.name = 'Arial'; gap.font.size = Pt(12); gap.bold = False
                add_inline(p, event['text'])
            elif event['family'] == 'table':
                add_table(doc, event['key'], model, warnings, information, layouts)
            else:
                key = event['key']
                add_figure(doc, key, model['figures'][key], model['figure_captions'][key], process, warnings)
        docs[output/f'{name}_Manuscript_v{version}.docx'] = stable_save(doc, output)
    if kind in ('both', 'supp'):
        doc = new_document(); add_heading(doc, 'Supplementary material', 16)
        for key in sorted(model['supp_figures'], key=lambda x: int(x[1:])):
            if doc.element.body[-2].tag != qn('w:p') or doc.paragraphs[-1].text: add_blank(doc)
            add_figure(doc, key, model['supp_figures'][key], model['supp_captions'][key], process, warnings)
            add_blank(doc)
        for key in sorted((x for x in model['tables'] if x.startswith('S')), key=lambda x: int(x[1:])):
            if doc.element.body[-2].tag != qn('w:p') or doc.paragraphs[-1].text: add_blank(doc)
            add_table(doc, key, model, warnings, information, layouts); add_blank(doc)
        docs[output/f'{name}_Supplementary_v{version}.docx'] = stable_save(doc, output)
    # Build both in memory before publishing either.
    for path, data in docs.items(): archive_write(path, data, root, day)
    manifest = dict(version=version, kind=kind, source=str(root), output=str(output),
                    figures=list(model['figures']), supp_figures=list(model['supp_figures']),
                    tables=list(model['tables']), citation_candidates=model['references'], table_layouts=layouts, warnings=warnings, information=information,
                    outputs={p.name: sha(p) for p in docs})
    (process/'assembly_manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    write_process_lists(process, model['references'], layouts)
    (process/'作者裁决清单.md').write_text('# Finalize 作者裁决清单\n\n'+
        '\n'.join('- '+w for w in warnings+information)+'\n\n沿用 PROJECT.md 待确认状态、冲突台账和各节核查说明中的未决事项。未作自动裁决。\n', encoding='utf-8')
    if update_tracking and kind == 'both': tracking(root, version, day)
    return manifest


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--project-root', required=True, type=Path)
    parser.add_argument('--kind', choices=('both', 'main', 'supp'), default='both')
    parser.add_argument('--version', type=int, default=1)
    parser.add_argument('--output-dir', type=Path)
    parser.add_argument('--name'); parser.add_argument('--date')
    parser.add_argument('--no-update-tracking', action='store_true')
    args = parser.parse_args()
    try:
        result = build(args.project_root, args.version, args.kind, args.output_dir, args.date,
                       args.name, not args.no_update_tracking)
        print(json.dumps(result, ensure_ascii=False, indent=2)); return 0
    except (ValueError, OSError) as e:
        print(json.dumps({'errors': [str(e)]}, ensure_ascii=False)); return 1


if __name__ == '__main__':
    raise SystemExit(main())
