"""Offline acceptance specification regression tests. All generated files stay in the explicit _tmp work dir."""
from __future__ import annotations
import argparse,copy,io,json,sys,unittest,uuid,zipfile
from pathlib import Path
from unittest.mock import patch
from lxml import etree as E
from docx import Document
from docx.shared import Pt
from PIL import Image
from extract_citations import q,extract,fields,package,xml_bytes,load,save,sha,scan_markers
from inject_zotero_fields import candidate,apply_candidate,gate
from check_references import check
from zotero_collect import Local,StopWrite,MetadataConflict,match_item,metadata,bibtex,write_one,verify_local_tables
WORK=None
BASELINE=Path(__file__).resolve().parents[1]/'assets/baseline-sample.docx'
class Regression(unittest.TestCase):
    def setUp(self):
        self.d=WORK/('t-'+uuid.uuid4().hex[:8]);self.d.mkdir(parents=True)
        self.source=self.d/'Test_Manuscript_v1.docx';self.cand=self.d/'process/candidate.docx';self.doc=Document()
        for t in ('[Manuscript title]','[Authors]','[Affiliations]','[Corresponding author]'):self.doc.add_paragraph(t)
        self.doc.add_paragraph('Introduction')
        p=self.doc.add_paragraph('Before ')
        r=p.add_run('(PM');r.bold=True;r.font.size=Pt(13)
        r=p.add_run('ID: 1)');r.italic=True
        p.add_run(' middle (PMID: 2) end (PMID: 1; PMID: 2).')
        self.doc.add_paragraph('Repeat (PMID: 1); manual (URL: https://example.org).')
        self.doc.add_paragraph('Figure 1. Caption (PMID: 2).')
        table=self.doc.add_table(rows=1,cols=1);table.cell(0,0).text='Table note (DOI: 10.1234/a).'
        self.doc.add_paragraph('Preprint (arXiv: 1512.03385).')
        image=self.d/'image.png';Image.new('RGB',(4,4),'blue').save(image);self.doc.add_picture(str(image))
        self.doc.add_paragraph('Discussion');self.doc.add_paragraph('Unchanged conclusion.');self.doc.save(self.source)
        self.man=extract(self.source);self.snap={'items':[]}
        for k,i,doi,pmid in [('AAA11111',101,'10.1234/a','1'),('BBB22222',102,'10.1234/b','2'),('CCC33333',103,'','')]:
            raw={'itemType':'journalArticle' if pmid else 'preprint','title':'Synthetic '+k,'DOI':doi,'extra':'PMID: '+pmid if pmid else 'arXiv: 1512.03385','url':''}
            csl={'id':i,'type':'article-journal' if pmid else 'article','title':raw['title'],'DOI':doi,'author':[{'family':'Test','given':'A'}],'issued':{'date-parts':[[2020]]}}
            self.snap['items'].append({'key':k,'raw':raw,'citation_item':{'id':i,'uris':['http://zotero.org/users/999/items/'+k],'itemData':csl},'resolved_identifiers':['arXiv: 1512.03385'] if not pmid else ['PMID: '+pmid]})
        for r in self.man['references']:
            if r['type']=='URL':continue
            r.update(status='既有',key='AAA11111' if r['id'] in ('PMID: 1','DOI: 10.1234/a') else 'BBB22222' if r['id']=='PMID: 2' else 'CCC33333')
    def inject(self):return candidate(self.source,self.man,self.snap,BASELINE,self.cand)
    def verify(self,man=None,snap=None):return check(self.source,self.cand,man or self.man,snap or self.snap,BASELINE)
    def tamper(self,fn,part='word/document.xml'):
        p=package(self.cand);r=E.fromstring(p[part]);fn(r);p[part]=xml_bytes(r)
        with zipfile.ZipFile(self.cand,'w',zipfile.ZIP_DEFLATED) as z:
            for n,data in p.items():z.writestr(n,data)
    def test_cross_run_multiple_repeated_grouped_caption_table_url_arxiv(self):
        self.assertEqual(len(self.man['occurrences']),8)
        self.assertEqual(len(self.man['references']),5)
        self.assertEqual(len(self.man['occurrences'][2]['identifiers']),2)
        self.assertTrue(any(o['kind']=='figure' for o in self.man['occurrences']))
        self.assertTrue(any(o['kind']=='table' for o in self.man['occurrences']))
        self.inject();report=self.verify();self.assertEqual(report['errors'],[],report)
        self.assertEqual(report['invariants']['I8']['status'],'WARN')
    def test_source_not_modified(self):
        before=sha(self.source);self.inject();self.assertEqual(before,sha(self.source))
    def test_baseline_split_instruction_parsing(self):
        fs=fields(E.fromstring(package(BASELINE)['word/document.xml']));self.assertEqual([f['kind'] for f in fs],['citation','citation','bibliography']);self.assertEqual(len(fs[0]['data']['citationItems']),2)
    def test_rerun_no_duplicate_fields(self):
        self.inject();nextman=extract(self.cand,previous=self.man)
        self.assertTrue(nextman['source_consistent'])
        for r,old in zip(nextman['references'],self.man['references']):r.update(status=old['status'],key=old.get('key'),reason=old.get('reason',''))
        second=self.d/'rerun/candidate.docx';candidate(self.cand,nextman,self.snap,BASELINE,second,self.man)
        report=check(self.cand,second,nextman,self.snap,BASELINE,self.man);self.assertEqual(report['errors'],[],report)
        self.assertEqual(len([f for f in fields(E.fromstring(package(second)['word/document.xml'])) if f['kind']=='citation']),7)
    def test_i9_manifest_drift_blocks(self):
        self.inject();prev=copy.deepcopy(self.man);prev['occurrences'][0]['marker']='(PMID: 999)'
        with self.assertRaisesRegex(ValueError,'I9'):extract(self.cand,previous=prev)
    def test_i1_missing_record(self):
        self.inject();bad=copy.deepcopy(self.man);bad['references'].pop();self.assertTrue(self.verify(man=bad)['errors'])
    def test_i1_illegal_status(self):
        self.inject();bad=copy.deepcopy(self.man);bad['references'][0]['status']='accepted';self.assertIn('I1',self.verify(man=bad)['errors'])
    def test_i2_empty_field(self):
        self.inject()
        def f(r):
            n=next(r.iter(q('instrText')));d=json.loads(n.text.split('CSL_CITATION',1)[1]);d['citationItems']=[];n.text=' ADDIN ZOTERO_ITEM CSL_CITATION '+json.dumps(d)+' '
        self.tamper(f);self.assertIn('I2',self.verify()['errors'])
    def test_i3_moved_field_same_paragraph(self):
        self.inject()
        def f(r):
            field=fields(r)[0];p=field['begin'].getparent()
            for n in field['runs']:p.remove(n);p.append(n)
        self.tamper(f);self.assertIn('I3',self.verify()['errors'])
    def test_i4_changed_payload(self):
        self.inject()
        def f(r):
            n=next(r.iter(q('instrText')));d=json.loads(n.text.split('CSL_CITATION',1)[1]);d['citationItems'][0]['itemData']['title']='Wrong';n.text=' ADDIN ZOTERO_ITEM CSL_CITATION '+json.dumps(d)+' '
        self.tamper(f);self.assertIn('I4',self.verify()['errors'])
    def test_i4_group_order(self):
        self.inject()
        def f(r):
            for n in r.iter(q('instrText')):
                if 'CSL_CITATION' not in (n.text or ''):continue
                d=json.loads(n.text.split('CSL_CITATION',1)[1])
                if len(d['citationItems'])>1:d['citationItems'].reverse();n.text=' ADDIN ZOTERO_ITEM CSL_CITATION '+json.dumps(d)+' ';break
        self.tamper(f);self.assertIn('I4',self.verify()['errors'])
    def test_i5_text_tamper(self):
        self.inject();self.tamper(lambda r:setattr(next(r.iter(q('t'))),'text','Changed'));self.assertIn('I5',self.verify()['errors'])
    def test_i5_direct_format_tamper(self):
        self.inject()
        def f(r):
            run=next(r.iter(q('r')));pr=run.find(q('rPr'))
            if pr is None:pr=E.SubElement(run,q('rPr'))
            E.SubElement(pr,q('i'))
        self.tamper(f);self.assertIn('I5',self.verify()['errors'])
    def test_i6_missing_registration(self):
        self.inject();self.tamper(lambda r:[r.remove(n) for n in list(r) if n.get('PartName')=='/docProps/custom.xml'],'[Content_Types].xml');self.assertIn('I6',self.verify()['errors'])
    def test_i6_wrong_style(self):
        self.inject()
        def f(r):
            for n in r.iter():
                if n.text:n.text=n.text.replace('elsevier-vancouver-short-author-list','wrong-style')
        self.tamper(f,'docProps/custom.xml');self.assertIn('I6',self.verify()['errors'])
    def test_i7_page_setup(self):
        self.inject();self.tamper(lambda r:next(r.iter(q('pgSz'))).set(q('w'),'42'));self.assertIn('I7',self.verify()['errors'])
    def test_i7_table_geometry(self):
        self.inject();self.tamper(lambda r:next(r.iter(q('gridCol'))).set(q('w'),'42'));self.assertIn('I7',self.verify()['errors'])
    def test_i7_image_bytes(self):
        self.inject();p=package(self.cand);name=next(n for n in p if n.startswith('word/media/'));p[name]+=b'changed'
        with zipfile.ZipFile(self.cand,'w') as z:
            for n,b in p.items():z.writestr(n,b)
        self.assertIn('I7',self.verify()['errors'])
    def test_i8_missing_reason(self):
        self.man['references'][-2 if self.man['references'][-2]['type']=='URL' else next(i for i,r in enumerate(self.man['references']) if r['type']=='URL')]['reason']=''
        with self.assertRaisesRegex(ValueError,'reason'):self.inject()
    def test_candidate_collision_refused(self):
        self.inject()
        with self.assertRaisesRegex(ValueError,'exists'):self.inject()
    def test_missing_baseline_refused(self):
        with self.assertRaises(OSError):candidate(self.source,self.man,self.snap,self.d/'absent.docx',self.cand)
    def test_source_hash_drift_refused(self):
        self.man['source_sha256']='0'*64
        with self.assertRaisesRegex(ValueError,'changed'):self.inject()
    def test_stage11_gate(self):
        (self.d/'PROJECT.md').write_text('| 11 | Finalize | x | 待确认 | x |\n',encoding='utf-8')
        with self.assertRaisesRegex(ValueError,'11'):gate(self.d)
    def test_stage11_written_authorization(self):
        (self.d/'PROJECT.md').write_text('| 11 | Finalize | x | 待确认 | x |\n',encoding='utf-8');a=self.d/'approval.md';a.write_text('Author explicitly permits this fixture project.',encoding='utf-8');gate(self.d,a)
    def test_apply_approval_required(self):
        self.inject()
        with self.assertRaisesRegex(ValueError,'approval'):apply_candidate(self.source,self.cand,self.man,self.snap,BASELINE,output_dir=self.d)
    def test_apply_check_blocks(self):
        self.inject();self.tamper(lambda r:setattr(next(r.iter(q('t'))),'text','Bad'))
        with self.assertRaisesRegex(ValueError,'ERROR'):apply_candidate(self.source,self.cand,self.man,self.snap,BASELINE,output_dir=self.d,approved=True)
    def test_atomic_archive_failure_rollback(self):
        self.inject();old=sha(self.source)
        with self.assertRaises(OSError):apply_candidate(self.source,self.cand,self.man,self.snap,BASELINE,output_dir=self.d,approved=True,fail_after='archive')
        self.assertEqual(sha(self.source),old);self.assertFalse((self.d/'Test_Manuscript_v2.docx').exists())
    def test_apply_success_archived_and_pending(self):
        self.inject();r=apply_candidate(self.source,self.cand,self.man,self.snap,BASELINE,output_dir=self.d,approved=True)
        self.assertEqual(r['status'],'待确认');self.assertFalse(self.source.exists());self.assertTrue(Path(r['new_manuscript']).exists());self.assertTrue((Path(r['archive'])/'05_final'/self.source.name).exists())
    def test_project_tracking_failure_rollback(self):
        project=self.d/'project';(project/'05_final').mkdir(parents=True);src=project/'05_final'/self.source.name;src.write_bytes(self.source.read_bytes());self.source=src;self.man['source_sha256']=sha(src);self.man['source']=str(src)
        p=project/'PROJECT.md';p.write_text('| 11 | Finalize | x | 已确认 | x |\n',encoding='utf-8');old=p.read_bytes();self.inject()
        with self.assertRaises(OSError):apply_candidate(src,self.cand,self.man,self.snap,BASELINE,project=project,approved=True,fail_after='tracking')
        self.assertEqual(p.read_bytes(),old);self.assertTrue(src.exists());self.assertFalse((project/'05_final/Test_Manuscript_v2.docx').exists())
    def test_doi_normalization_existing(self):
        rows=[{'key':'X','data':{'DOI':'https://doi.org/10.1234/A.'}}];found,method=match_item(rows,{'id':'DOI: 10.1234/a','type':'DOI','value':'10.1234/a'});self.assertEqual(found['key'],'X');self.assertEqual(method,'DOI')
    def test_pmid_resolves_existing_doi_without_pmid(self):
        rows=[{'key':'X','data':{'DOI':'10.1234/a'}}];found,method=match_item(rows,{'id':'PMID: 1','type':'PMID','value':'1'},{'data':{'DOI':'10.1234/a','extra':'PMID: 1'}});self.assertEqual(method,'DOI')
    def test_title_year_fallback(self):
        rows=[{'key':'X','data':{'title':'Test: A B','date':'2020-01'}}];_,method=match_item(rows,{'id':'PMID: 1','type':'PMID','value':'1'},{'data':{'title':'test a-b','date':'2020'}});self.assertEqual(method,'title-year')
    def test_duplicate_matches_stop(self):
        rows=[{'key':k,'data':{'DOI':'10.1234/a'}} for k in ('X','Y')]
        with self.assertRaises(MetadataConflict):match_item(rows,{'id':'DOI: 10.1234/a','type':'DOI','value':'10.1234/a'})
    def test_probe_three_failures_and_no_process_control(self):
        client=Local();calls=[]
        def fail(*args,**kwargs):calls.append(args);raise OSError('offline')
        with patch.object(client,'request',side_effect=fail):r=client.probe(attempts=3,delay=0)
        self.assertFalse(r['ready']);self.assertEqual(len(calls),3)
    def test_b_wrong_target_no_write(self):
        class Fake:
            def items(self):return []
            def selected(self):return {'name':'WRONG','editable':True,'id':12}
            def request(self,*a,**k):raise AssertionError('Write must not happen')
        with self.assertRaisesRegex(StopWrite,'target mismatch'):write_one(Fake(),{'DOI':'10.1234/a'}, {'key':'T','data':{'name':'_skill-test-20261016'}},'B',self.d,{'id':'DOI: 10.1234/a','type':'DOI','value':'10.1234/a'})
    def test_existing_item_no_write(self):
        class Fake:
            def items(self):return [{'key':'X','data':{'DOI':'10.1234/a'}}]
            def request(self,*a,**k):raise AssertionError('No duplicate write')
        item,added,_=write_one(Fake(),{'DOI':'10.1234/a'},{'key':'T','data':{'name':'test'}},'A',self.d,{'id':'DOI: 10.1234/a','type':'DOI','value':'10.1234/a'});self.assertFalse(added)
    def test_pubmed_structured_full_authors(self):
        xml=b'<PubmedArticleSet><PubmedArticle><MedlineCitation><Article><ArticleTitle>Public test</ArticleTitle><Journal><Title>Test Journal</Title><ISOAbbreviation>Test J</ISOAbbreviation><JournalIssue><PubDate><Year>2020</Year></PubDate></JournalIssue></Journal><Pagination><StartPage>1122</StartPage><EndPage>1131.e9</EndPage><MedlinePgn>1122-1131.e9</MedlinePgn></Pagination><AuthorList><Author><LastName>van Test</LastName><ForeName>A B</ForeName></Author><Author><CollectiveName>Study Group</CollectiveName></Author></AuthorList></Article></MedlineCitation><PubmedData><ArticleIdList><ArticleId IdType="doi">10.1234/a</ArticleId></ArticleIdList></PubmedData></PubmedArticle></PubmedArticleSet>'
        with patch('zotero_collect.public_get',return_value=xml):m=metadata({'type':'PMID','value':'1','id':'PMID: 1'})
        self.assertEqual(m['data']['creators'][0]['lastName'],'van Test');self.assertEqual(m['data']['creators'][1]['name'],'Study Group');self.assertEqual(m['data']['extra'],'PMID: 1');self.assertEqual(m['data']['pages'],'1122-1131.e9')
    def test_native_pmid_dedup_without_doi(self):
        found,method=match_item([{'key':'X','data':{'PMID':'1'}}],{'id':'PMID: 1','type':'PMID','value':'1'})
        self.assertEqual(found['key'],'X');self.assertEqual(method,'PMID')
    def test_local_verification_metadata_conflict_stops(self):
        folder=self.d/'04_manuscript/Introduction';folder.mkdir(parents=True)
        (folder/'Introduction_核查说明.md').write_text('| 标识 | 题名 | 第一作者 年份 | 期刊 |\n|---|---|---|---|\n| PMID: 1 | Wrong title | Test 2020 | Journal |\n',encoding='utf-8')
        with self.assertRaises(MetadataConflict):verify_local_tables(self.d,{'id':'PMID: 1'},{'data':{'title':'Correct title','date':'2020','publicationTitle':'Journal','creators':[{'lastName':'Test'}]}})
    def test_finalize_candidate_metadata_conflict_stops(self):
        folder=self.d/'_process/Finalize/day';folder.mkdir(parents=True)
        save(folder/'assembly_manifest.json',{'citation_candidates':[{'id':'PMID: 1','fields':{'题名':'Wrong title'}}]})
        with self.assertRaises(MetadataConflict):verify_local_tables(self.d,{'id':'PMID: 1'},{'data':{'title':'Correct title'}})
    def test_i9_field_item_identity_drift(self):
        self.inject()
        def f(root):
            n=next(root.iter(q('instrText')));d=json.loads(n.text.split('CSL_CITATION',1)[1]);d['citationItems'][0]['uris']=['http://zotero.org/users/999/items/WRONG'];n.text=' ADDIN ZOTERO_ITEM CSL_CITATION '+json.dumps(d)+' '
        self.tamper(f)
        with self.assertRaisesRegex(ValueError,'I9'):extract(self.cand,previous=self.man)
    def test_http403_stops_without_connector_fallback(self):
        import urllib.error
        client=Local();client.auth={'key':'SYNTHETIC'};client.server='TEST'
        error=urllib.error.HTTPError('http://127.0.0.1:23119/api/users/0/items',403,'denied',{},io.BytesIO(b'denied'))
        with patch.object(client.opener,'open',side_effect=error) as opened:
            with self.assertRaisesRegex(StopWrite,'403'):client.request('/api/users/0/items','POST',[],write=True)
        self.assertEqual(opened.call_count,1)
    def test_write_timeout_does_not_repost(self):
        class Fake:
            calls=0
            def items(self):return []
            def request(self,*a,**k):self.calls+=1;raise OSError('lost response')
        f=Fake()
        with self.assertRaises(OSError):write_one(f,{'DOI':'10.1234/a'},{'key':'T','data':{'name':'test'}},'A',self.d,{'id':'DOI: 10.1234/a','type':'DOI','value':'10.1234/a'})
        self.assertEqual(f.calls,1)
    def test_malformed_marker_not_silently_dropped(self):
        with self.assertRaisesRegex(ValueError,'Malformed'):scan_markers('Text (PMID: unknown).')
        with self.assertRaisesRegex(ValueError,'DOI'):scan_markers('Text (DOI: 12345).')
    def test_apply_writes_inside_references_section(self):
        project=self.d/'project';(project/'05_final').mkdir(parents=True);src=project/'05_final'/self.source.name;src.write_bytes(self.source.read_bytes());self.source=src;self.man['source_sha256']=sha(src);self.man['source']=str(src)
        (project/'PROJECT.md').write_text('| 11 | Finalize | x | 已确认 | x |\n',encoding='utf-8');folder=project/'04_manuscript';folder.mkdir();decision=folder/'作者决定.md';decision.write_text('# 作者决定\n\n## References\n\nOld decision.\n\n## Another\n\nKeep here.\n',encoding='utf-8');self.inject()
        apply_candidate(src,self.cand,self.man,self.snap,BASELINE,project=project,approved=True)
        text=decision.read_text(encoding='utf-8');self.assertLess(text.index('字段版主稿'),text.index('## Another'));self.assertIn('Keep here.',text)
    def test_collect_offline_cli_reports_pending_exit2(self):
        import zotero_collect
        man=self.d/'manifest.json';save(man,self.man)
        args=['zotero_collect.py','collect','--manifest',str(man),'--output-dir',str(self.d/'pending'),'--collection','test']
        with patch.object(sys,'argv',args),patch.object(Local,'probe',return_value={'ready':False,'attempts':[1,2,3]}),patch.object(Local,'items',side_effect=AssertionError('No read/write while offline')):
            self.assertEqual(zotero_collect.main(),2)
        self.assertEqual(load(self.d/'pending/zotero写入报告.json')['status'],'pending')
    def test_preprint_second_sample_missing_warns_without_block(self):
        b=self.d/'baseline-sample.docx';b.write_bytes(BASELINE.read_bytes())
        receipt=candidate(self.source,self.man,self.snap,b,self.cand)
        self.assertTrue(receipt['warnings']);report=check(self.source,self.cand,self.man,self.snap,b);self.assertEqual(report['errors'],[]);self.assertTrue(any('preprint' in w for w in report['warnings']))
    def test_arxiv_literal_names_without_invented_splitting(self):
        atom=b'<feed xmlns="http://www.w3.org/2005/Atom"><entry><title>Public test</title><published>2015-12-10T00:00:00Z</published><author><name>First Middle van Last</name></author></entry></feed>'
        with patch('zotero_collect.public_get',return_value=atom):m=metadata({'type':'arXiv','value':'1512.03385','id':'arXiv: 1512.03385'})
        self.assertEqual(m['data']['itemType'],'preprint');self.assertEqual(m['data']['creators'],[{'creatorType':'author','name':'First Middle van Last'}]);self.assertTrue(m['warnings'])
    def journal_fixture(self,local,kind='table',**changes):
        root=self.d/kind
        fields={'题名':'Public title','期刊':local,'第一作者 年份':'Test 2020'}
        if kind=='table':
            folder=root/'04_manuscript/Introduction';folder.mkdir(parents=True,exist_ok=True)
            (folder/'Introduction_核查说明.md').write_text('| 标识 | 题名 | 期刊 | 第一作者 年份 |\n|---|---|---|---|\n| PMID: 1 | '+fields['题名']+' | '+local+' | '+fields['第一作者 年份']+' |\n',encoding='utf-8')
        else:
            folder=root/'_process/Finalize/day';folder.mkdir(parents=True,exist_ok=True);save(folder/'assembly_manifest.json',{'citation_candidates':[{'id':'PMID: 1','fields':fields}]})
        data={'title':'Public title','publicationTitle':'Journal of Machine Learning Research','journalAbbreviation':'J Mach Learn Res','date':'2020','creators':[{'lastName':'Test','firstName':'A'}]};data.update(changes)
        return root,{'data':data,'warnings':[]}
    def test_journal_iso_abbreviation_accepted_warn_table(self):
        root,meta=self.journal_fixture('J Mach Learn Res')
        verify_local_tables(root,{'id':'PMID: 1'},meta)
        self.assertTrue(meta['warnings']);self.assertIn('J Mach Learn Res',str(meta['warnings']));self.assertIn('Journal of Machine Learning Research',str(meta['warnings']))
    def test_journal_iso_abbreviation_accepted_warn_manifest(self):
        root,meta=self.journal_fixture('J. Mach. Learn. Res.','manifest')
        verify_local_tables(root,{'id':'PMID: 1'},meta);self.assertTrue(meta['warnings'])
    def test_journal_full_name_still_accepted(self):
        for kind in ('table','manifest'):
            root,meta=self.journal_fixture('Journal of Machine Learning Research',kind)
            verify_local_tables(root,{'id':'PMID: 1'},meta)
    def test_journal_different_title_rejected_both_sources(self):
        for kind in ('table','manifest'):
            with self.subTest(kind=kind):
                root,meta=self.journal_fixture('Journal of Clinical Medicine',kind)
                with self.assertRaises(MetadataConflict):verify_local_tables(root,{'id':'PMID: 1'},meta)
    def test_journal_equivalence_does_not_hide_year_conflict(self):
        for kind in ('table','manifest'):
            with self.subTest(kind=kind):
                root,meta=self.journal_fixture('J Mach Learn Res',kind,date='2021')
                with self.assertRaises(MetadataConflict):verify_local_tables(root,{'id':'PMID: 1'},meta)
    def test_journal_equivalence_does_not_hide_author_conflict(self):
        for kind in ('table','manifest'):
            with self.subTest(kind=kind):
                root,meta=self.journal_fixture('J Mach Learn Res',kind,creators=[{'lastName':'Other'}])
                with self.assertRaises(MetadataConflict):verify_local_tables(root,{'id':'PMID: 1'},meta)
    def add_mixed_run(self,placement='leading',text='Before (PMID: 1) after',tag='lastRenderedPageBreak'):
        p=E.Element(q('p'));run=E.SubElement(p,q('r'));run.set(q('rsidR'),'00ABCDEF');pr=E.SubElement(run,q('rPr'));E.SubElement(pr,q('b'));E.SubElement(pr,q('sz')).set(q('val'),'26')
        marker=E.Element(q(tag));marker.set('testAttribute','keep')
        t=E.Element(q('t'));t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');t.text=text
        run.extend([marker,t] if placement=='leading' else [t,marker])
        parts=package(self.source);root=E.fromstring(parts['word/document.xml']);body=root.find(q('body'));body.insert(len(body)-1,p);parts['word/document.xml']=xml_bytes(root)
        with zipfile.ZipFile(self.source,'w') as z:
            for n,b in parts.items():z.writestr(n,b)
        old={r['id']:r for r in self.man['references']};self.man=extract(self.source)
        for ref in self.man['references']:ref.update(old[ref['id']])
        return marker
    def independent_pagination(self,path):
        result=[]
        for pi,p in enumerate(E.fromstring(package(path)['word/document.xml']).iter(q('p'))):
            offset=0
            for node in p.iter():
                if next(node.iterancestors(q('p')),None) is not p:continue
                if node.tag==q('lastRenderedPageBreak'):
                    r=node.getparent();pr=r.find(q('rPr'));result.append((pi,offset,dict(node.attrib),dict(r.attrib),E.tostring(pr,method='c14n') if pr is not None else b''))
                elif node.tag==q('t'):offset+=len(node.text or '')
                elif node.tag in (q('tab'),q('br'),q('cr')):offset+=1
        return result
    def mixed_positive(self,placement,text='Before (PMID: 1) after'):
        self.add_mixed_run(placement,text);before=self.independent_pagination(self.source);self.inject();report=self.verify()
        self.assertEqual(report['errors'],[],report);self.assertEqual(report['invariants'].get('I7a',{}).get('status'),'PASS');self.assertEqual(before,self.independent_pagination(self.cand))
    def test_mixed_leading_page_marker_split(self):self.mixed_positive('leading')
    def test_mixed_trailing_page_marker_split(self):self.mixed_positive('trailing')
    def test_mixed_page_marker_at_citation_start(self):self.mixed_positive('leading','(PMID: 1) after')
    def test_mixed_page_marker_at_citation_end(self):self.mixed_positive('trailing','Before (PMID: 1)')
    def test_mixed_page_marker_with_citation_entire_run(self):self.mixed_positive('leading','(PMID: 1)')
    def test_i7a_moved_page_marker_rejected(self):
        self.add_mixed_run();self.inject()
        def move(root):
            m=next(root.iter(q('lastRenderedPageBreak')));run=m.getparent();run.remove(m);run.append(m)
        self.tamper(move);report=self.verify()
        self.assertEqual(report['invariants']['I5']['status'],'PASS');self.assertEqual(report['invariants']['I7']['status'],'PASS');self.assertIn('I7a',report['errors'])
    def test_i7a_page_marker_attribute_change_rejected(self):
        self.add_mixed_run();self.inject();self.tamper(lambda root:next(root.iter(q('lastRenderedPageBreak'))).set('testAttribute','changed'))
        self.assertIn('I7a',self.verify()['errors'])
    def test_unproven_soft_hyphen_run_still_rejected(self):
        self.add_mixed_run(tag='softHyphen')
        with self.assertRaisesRegex(ValueError,'unsupported|non-text'):self.inject()
    def test_unproven_no_break_hyphen_run_still_rejected(self):
        self.add_mixed_run(tag='noBreakHyphen')
        with self.assertRaisesRegex(ValueError,'unsupported|non-text'):self.inject()
    def test_unproven_internal_page_marker_run_still_rejected(self):
        parts=package(self.source);root=E.fromstring(parts['word/document.xml']);p=E.SubElement(root.find(q('body')),q('p'));run=E.SubElement(p,q('r'));E.SubElement(run,q('t')).text='Before ';E.SubElement(run,q('lastRenderedPageBreak'));E.SubElement(run,q('t')).text='(PMID: 1) after'
        parts['word/document.xml']=xml_bytes(root)
        with zipfile.ZipFile(self.source,'w') as z:
            for n,b in parts.items():z.writestr(n,b)
        old={r['id']:r for r in self.man['references']};self.man=extract(self.source)
        for ref in self.man['references']:ref.update(old[ref['id']])
        with self.assertRaisesRegex(ValueError,'unsupported|non-text'):self.inject()
    def url_collect_fixture(self,entry=None,mapping_key='URL: https://example.org',map_note='Author approved this cited version',prefix_missing_item=False):
        import zotero_collect as z
        doc=Document()
        if prefix_missing_item:doc.add_paragraph('(PMID: 999)')
        doc.add_paragraph('Before (Author, 2020; URL: https://example.org) after')
        doc.save(self.source);self.man=extract(self.source)
        item=self.snap['items'][0];raw=copy.deepcopy(item['raw']);raw['collections']=['LZPSVVK8'];existing={'key':item['key'],'data':raw}
        csl=copy.deepcopy(item['citation_item']['itemData']);csl['id']=item['citation_item']['uris'][0]
        class Fake:
            def probe(fake):return {'ready':True}
            def items(fake):return [copy.deepcopy(existing)]
            def all(fake,path):return [{'key':'LZPSVVK8','data':{'name':'test','parentCollection':False}}]
            def request(fake,path,*args,**kwargs):
                if path=='/api/users/0/items/'+item['key']+'?format=csljson':return copy.deepcopy(csl),200,{}
                raise AssertionError('Unexpected request; no writes allowed: '+path)
        out=self.d/('collect-'+uuid.uuid4().hex[:6]);mapping=None
        if entry is not None:
            mapping=self.d/'resolve.json';save(mapping,{'note':map_note,'collection_key':'LZPSVVK8','mapping':{mapping_key:entry}})
        return z,Fake(),out,mapping,{'AAA11111':101,'__collection_ids__':{'LZPSVVK8':25}}
    def collect_url(self,entry=None,mapping_key='URL: https://example.org'):
        z,client,out,mapping,ids=self.url_collect_fixture(entry,mapping_key)
        with patch.object(z,'Local',return_value=client),patch.object(z,'integer_ids',return_value=ids),patch.object(z,'metadata',side_effect=AssertionError('URL metadata must not be fetched')),patch.object(z,'write_one',side_effect=AssertionError('Author mappings must not write')):
            kw={'resolve_file':mapping} if mapping else {}
            report=z.collect(self.man,out,'test',channel='A',work_dir=self.d/'_tmp',**kw)
        self.snap=load(out/'zotero条目快照.json');return report,out,mapping
    def test_resolve_file_existing_key_injects_with_provenance(self):
        report,out,mapping=self.collect_url({'item_key':'AAA11111','note':'Specific source version'})
        ref=self.man['references'][0];self.assertEqual(ref['status'],'作者指定');self.assertEqual(ref['resolution']['source'],'作者映射');self.assertEqual(ref['resolution']['file'],str(mapping.resolve()));self.assertEqual(ref['resolution']['item_key'],'AAA11111');self.assertIn('Specific source version',ref['resolution']['note']);self.assertIn('Author approved',ref['resolution']['note']);self.assertEqual(report['counts']['作者指定'],1);self.assertEqual(report['added_keys'],[])
        self.inject();checked=self.verify();self.assertEqual(checked['errors'],[],checked);self.assertTrue(any('作者映射' in w for w in checked['warnings']));self.assertEqual(checked['author_resolutions'][0]['resolution']['item_key'],'AAA11111')
        before=Document(self.source).paragraphs[0].text;after=Document(self.cand).paragraphs[0].text;self.assertEqual(before,after)
        field=next(f for f in fields(E.fromstring(package(self.cand)['word/document.xml'])) if f['kind']=='citation');self.assertEqual(field['data']['properties']['plainCitation'],'URL: https://example.org')
        self.assertIn('## 已由作者映射',(out/'缺口与人工清单.md').read_text(encoding='utf-8'))
    def test_resolve_file_bare_url_compatibility_warns(self):
        report,out,mapping=self.collect_url({'item_key':'AAA11111'},'https://example.org')
        self.assertEqual(self.man['references'][0]['status'],'作者指定');self.assertTrue(any('前缀' in w for w in self.man['references'][0]['warnings']))
    def test_resolve_file_missing_key_hard_error_before_other_metadata(self):
        z,client,out,mapping,ids=self.url_collect_fixture({'item_key':'ZZZZZZZZ'},prefix_missing_item=True)
        with patch.object(z,'Local',return_value=client),patch.object(z,'integer_ids',return_value=ids),patch.object(z,'metadata',side_effect=AssertionError('Preflight must precede other metadata')),patch.object(z,'write_one',side_effect=AssertionError('No writes')):
            with self.assertRaisesRegex(StopWrite,'ZZZZZZZZ'):z.collect(self.man,out,'test',channel='A',resolve_file=mapping)
        self.assertEqual(load(out/'zotero写入报告.json')['status'],'stopped');self.assertFalse((out/'zotero条目快照.json').exists())
    def test_resolve_file_absent_preserves_legacy_url_behavior(self):
        report,out,_=self.collect_url()
        self.assertEqual(self.man['references'][0]['status'],'人工保留');self.assertEqual(self.man['references'][0]['reason'],'URL-only 不自动抓取');self.assertNotIn('resolution',self.man['references'][0]);self.assertEqual(report['counts'],{'既有':0,'新增':0,'人工保留':1});self.assertEqual(report['results'],[]);self.assertEqual(self.snap['items'],[])
        self.inject();self.assertEqual(self.verify()['errors'],[]);self.assertEqual(len([f for f in fields(E.fromstring(package(self.cand)['word/document.xml'])) if f['kind']=='citation']),0)
    def test_resolve_file_unmatched_url_keeps_manual(self):
        report,_,_=self.collect_url({'item_key':'AAA11111'},'URL: https://unrelated.org')
        self.assertEqual(self.man['references'][0]['status'],'人工保留');self.assertNotIn('resolution',self.man['references'][0])
    def test_resolve_file_case_mismatch_rejected(self):
        with self.assertRaisesRegex((StopWrite,ValueError),'大小写|匹配|mapping|Mapping'):self.collect_url({'item_key':'AAA11111'},'URL: https://EXAMPLE.org')
    def test_resolve_file_whitespace_mismatch_rejected(self):
        with self.assertRaisesRegex((StopWrite,ValueError),'空白|匹配|mapping|Mapping'):self.collect_url({'item_key':'AAA11111'},' URL: https://example.org')
    def test_resolve_file_alternate_existing_identifier(self):
        report,_,_=self.collect_url({'identifier':{'type':'DOI','value':'10.1234/a'}})
        self.assertEqual(self.man['references'][0]['status'],'作者指定');self.assertEqual(self.man['references'][0]['resolution']['item_key'],'AAA11111');self.assertEqual(report['added_keys'],[])
    def test_resolve_file_invalid_resolution_key_rejected_by_i4(self):
        self.collect_url({'item_key':'AAA11111'});self.inject();bad=copy.deepcopy(self.man);bad['references'][0]['resolution']['item_key']='BBB22222'
        self.assertIn('I4',self.verify(man=bad)['errors'])
    def test_resolve_file_field_key_tamper_rejected(self):
        self.collect_url({'item_key':'AAA11111'});self.inject()
        def tamper(root):
            n=next(root.iter(q('instrText')));data=json.loads(n.text.split('CSL_CITATION',1)[1]);data['citationItems'][0]['uris']=['http://zotero.org/users/999/items/BBB22222'];n.text=' ADDIN ZOTERO_ITEM CSL_CITATION '+json.dumps(data)+' '
        self.tamper(tamper);self.assertIn('I4',self.verify()['errors'])
    def test_resolve_file_field_source_rerun_is_idempotent(self):
        self.collect_url({'item_key':'AAA11111'});self.inject();current=extract(self.cand,previous=self.man)
        for ref,old in zip(current['references'],self.man['references']):ref.update(old)
        other=self.d/'rerun/fields.docx';candidate(self.cand,current,self.snap,BASELINE,other,self.man);report=check(self.cand,other,current,self.snap,BASELINE,self.man)
        self.assertEqual(report['errors'],[],report);self.assertEqual(len([f for f in fields(E.fromstring(package(other)['word/document.xml'])) if f['kind']=='citation']),1)
    def test_bibtex_pmid_and_full_authors(self):
        data={'itemType':'journalArticle','title':'T','creators':[{'lastName':'van Test','firstName':'A B'},{'name':'Study Group'}],'extra':'PMID: 1','date':'2020'};b=bibtex(data);self.assertIn('pmid = {1}',b);self.assertIn('van Test, A B and {Study Group}',b)

def main():
    global WORK,BASELINE
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--work-dir',required=True,type=Path);p.add_argument('--baseline',type=Path);p.add_argument('--report',type=Path);a=p.parse_args()
    if a.work_dir.name!='_tmp':raise ValueError('Use an explicit _tmp work dir')
    WORK=a.work_dir.resolve()/('r-'+uuid.uuid4().hex[:8]);WORK.mkdir(parents=True)
    if a.baseline:BASELINE=a.baseline
    suite=unittest.defaultTestLoader.loadTestsFromTestCase(Regression);buffer=io.StringIO();result=unittest.TextTestRunner(stream=buffer,verbosity=2).run(suite);print(buffer.getvalue())
    if a.report:save(a.report,{'tests':result.testsRun,'failures':len(result.failures),'errors':len(result.errors),'passed':result.wasSuccessful(),'work_dir':str(WORK),'log':buffer.getvalue()})
    return 0 if result.wasSuccessful() else 1
if __name__=='__main__':sys.exit(main())
